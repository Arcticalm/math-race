from __future__ import annotations

import argparse
import csv
import json
import math
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Iterable

import rasterio
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook
from rasterio.transform import rowcol


ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT / "data"
BASE_DATA = DATA / "无人机应急物资运输基础数据"
DEM_PATH = (
    DATA
    / "镇龙乡地理空间数据"
    / "镇龙乡及周边地理数据"
    / "数字高程模型数据（DEM）"
    / "镇龙乡及周边30米DEM.tif"
)


@dataclass(frozen=True)
class Aircraft:
    code: str
    name: str
    empty_mass_kg: float
    max_payload_kg: float
    volume_m3: float
    cruise_speed_mps: float
    empty_range_m: float
    full_range_m: float
    usable_energy_kwh: float
    reserve_fraction: float
    prep_s: float
    load_each_s: float
    handoff_base_s: float
    handoff_each_s: float
    climb_speed_mps: float
    descent_speed_mps: float
    climb_efficiency: float
    descent_efficiency: float


@dataclass(frozen=True)
class Node:
    code: str
    longitude: float
    latitude: float
    elevation_m: float


@dataclass(frozen=True)
class Box:
    code: str
    site: str
    item_type: str
    mass_kg: float
    volume_m3: float


@dataclass(frozen=True)
class Flight:
    energy_kwh: float
    time_s: float
    outbound_energy_kwh: float
    return_energy_kwh: float
    outbound_time_s: float
    return_time_s: float
    range_budget_m: float
    return_range_budget_m: float
    horizontal_distance_m: float
    cruise_altitude_m: float


@dataclass(frozen=True)
class Candidate:
    site: str
    aircraft: str
    boxes: tuple[str, ...]
    mass_kg: float
    volume_m3: float
    energy_kwh: float
    outbound_energy_kwh: float
    return_energy_kwh: float
    flight_time_s: float
    outbound_time_s: float
    return_time_s: float
    work_time_s: float
    range_budget_m: float
    return_range_budget_m: float
    reserve_kwh: float
    return_soc_percent: float
    feasible: bool


def _rows(path: Path, sheet: str | None = None) -> list[tuple]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[sheet] if sheet else workbook.active
    return [tuple(row) for row in worksheet.iter_rows(values_only=True)]


def load_aircraft() -> dict[str, Aircraft]:
    rows = _rows(BASE_DATA / "运输无人机数据.xlsx")
    result = {}
    for row in rows:
        if row[0] in {"A", "B", "C"} and isinstance(row[3], (int, float)):
            result[row[0]] = Aircraft(
                code=row[0], name=row[1], empty_mass_kg=float(row[2]),
                max_payload_kg=float(row[3]), volume_m3=float(row[4]),
                cruise_speed_mps=float(row[5]), empty_range_m=float(row[6]),
                full_range_m=float(row[7]), usable_energy_kwh=float(row[8]),
                reserve_fraction=float(row[9]) / 100, prep_s=float(row[10]),
                load_each_s=float(row[11]), handoff_base_s=float(row[12]),
                handoff_each_s=float(row[13]), climb_speed_mps=float(row[14]),
                descent_speed_mps=float(row[15]), climb_efficiency=float(row[16]),
                descent_efficiency=float(row[17]),
            )
    if len(result) != 3:
        raise ValueError(f"Expected three aircraft types, got {sorted(result)}")
    return result


def load_nodes() -> tuple[Node, dict[str, Node]]:
    rows = _rows(BASE_DATA / "调度中心与服务区.xlsx")
    node_rows = [row for row in rows if row and row[0] == "O01"]
    sites = [row for row in rows if row and isinstance(row[0], str) and row[0].startswith("S")]
    if len(node_rows) != 1 or len(sites) != 15:
        raise ValueError(f"Expected O01 and 15 service areas; got {len(node_rows)}, {len(sites)}")

    def convert(row: tuple) -> Node:
        return Node(row[0], float(row[2]), float(row[3]), float(row[4]))

    return convert(node_rows[0]), {row[0]: convert(row) for row in sites}


def load_boxes() -> dict[str, list[Box]]:
    rows = _rows(BASE_DATA / "物资需求与配送时限.xlsx", "逐箱货箱清单")
    result: dict[str, list[Box]] = {}
    seen: set[str] = set()
    for row in rows[1:]:
        if not row[0]:
            continue
        code = str(row[0])
        if code in seen:
            raise ValueError(f"Duplicate box identifier: {code}")
        seen.add(code)
        box = Box(code, str(row[1]), str(row[2]), float(row[3]), float(row[4]))
        result.setdefault(box.site, []).append(box)
    if len(seen) != 80 or len(result) != 15:
        raise ValueError(f"Expected 80 boxes across 15 sites; got {len(seen)} and {len(result)}")
    return result


def great_circle_distance_m(start: Node, end: Node) -> float:
    radius_m = 6_371_008.8
    lat1, lat2 = math.radians(start.latitude), math.radians(end.latitude)
    delta_lat = lat2 - lat1
    delta_lon = math.radians(end.longitude - start.longitude)
    term = math.sin(delta_lat / 2) ** 2 + math.cos(lat1) * math.cos(lat2) * math.sin(delta_lon / 2) ** 2
    return 2 * radius_m * math.asin(math.sqrt(term))


def sample_leg(dem: rasterio.io.DatasetReader, start: Node, end: Node) -> tuple[float, float]:
    """Return maximum sampled DEM elevation and great-circle horizontal distance."""
    distance = great_circle_distance_m(start, end)
    intervals = max(1, math.ceil(distance / 30))
    elevations = []
    for index in range(intervals + 1):
        fraction = index / intervals
        longitude = start.longitude + fraction * (end.longitude - start.longitude)
        latitude = start.latitude + fraction * (end.latitude - start.latitude)
        row, col = rowcol(dem.transform, longitude, latitude)
        if not (0 <= row < dem.height and 0 <= col < dem.width):
            raise ValueError(f"Route {start.code}-{end.code} leaves DEM extent")
        value = float(dem.read(1, window=((row, row + 1), (col, col + 1)))[0, 0])
        if dem.nodata is not None and math.isclose(value, dem.nodata):
            raise ValueError(f"Route {start.code}-{end.code} crosses DEM NoData")
        elevations.append(value)
    return max(elevations), distance


def load_energy_per_meter(aircrafts: dict[str, Aircraft]) -> dict[str, float]:
    """Calibrate cruise energy per meter from the supplied standard ranges.

    The statement does not prescribe a numerical horizontal-energy function.
    This explicitly documented baseline treats usable battery energy divided by
    standard range as the average horizontal energy rate.
    """
    return {
        code: aircraft.usable_energy_kwh / aircraft.empty_range_m
        for code, aircraft in aircrafts.items()
    }


def leg_flight(
    aircraft: Aircraft,
    energy_per_meter: float,
    distance_m: float,
    cruise_altitude_m: float,
    start_altitude_m: float,
    end_altitude_m: float,
    payload_kg: float,
    mass_override_kg: float | None = None,
) -> tuple[float, float, float]:
    climb_m = max(0.0, cruise_altitude_m - start_altitude_m)
    descent_m = max(0.0, cruise_altitude_m - end_altitude_m)
    flight_time_s = (
        climb_m / aircraft.climb_speed_mps
        + distance_m / aircraft.cruise_speed_mps
        + descent_m / aircraft.descent_speed_mps
    )
    payload_ratio = payload_kg / aircraft.max_payload_kg
    range_m = aircraft.empty_range_m - (aircraft.empty_range_m - aircraft.full_range_m) * payload_ratio**1.5
    horizontal_energy_kwh = energy_per_meter * distance_m * aircraft.empty_range_m / range_m
    flight_mass_kg = mass_override_kg if mass_override_kg is not None else aircraft.empty_mass_kg + payload_kg
    climb_energy_kwh = (
        flight_mass_kg * 9.80665 * climb_m
        / (3_600_000 * aircraft.climb_efficiency)
    )
    descent_energy_kwh = 0.0
    if aircraft.descent_efficiency > 0:
        descent_energy_kwh = (
            flight_mass_kg * 9.80665 * descent_m
            / (3_600_000 * aircraft.descent_efficiency)
        )
    return flight_time_s, horizontal_energy_kwh + climb_energy_kwh + descent_energy_kwh, range_m


def flight_for_payload(
    aircraft: Aircraft,
    energy_per_meter: float,
    base: Node,
    site: Node,
    maximum_ground_elevation: float,
    payload_kg: float,
) -> Flight:
    cruise_altitude = maximum_ground_elevation + 50
    base_altitude = base.elevation_m
    site_altitude = site.elevation_m + 30
    outbound_time, outbound_energy, range_m = leg_flight(
        aircraft, energy_per_meter, great_circle_distance_m(base, site),
        cruise_altitude, base_altitude, site_altitude, payload_kg,
    )
    return_time, return_energy, _ = leg_flight(
        aircraft, energy_per_meter, great_circle_distance_m(site, base),
        cruise_altitude, site_altitude, base_altitude, 0.0, aircraft.empty_mass_kg,
    )
    horizontal_round_trip = 2 * great_circle_distance_m(base, site)
    return Flight(
        energy_kwh=outbound_energy + return_energy,
        time_s=outbound_time + return_time,
        outbound_energy_kwh=outbound_energy,
        return_energy_kwh=return_energy,
        outbound_time_s=outbound_time,
        return_time_s=return_time,
        range_budget_m=range_m,
        return_range_budget_m=aircraft.empty_range_m,
        horizontal_distance_m=horizontal_round_trip,
        cruise_altitude_m=cruise_altitude,
    )


def max_safe_payload(
    aircraft: Aircraft,
    energy_per_meter: float,
    base: Node,
    site: Node,
    maximum_ground_elevation: float,
) -> tuple[float, Flight]:
    low, high = 0.0, aircraft.max_payload_kg
    best = flight_for_payload(aircraft, energy_per_meter, base, site, maximum_ground_elevation, low)
    if best.energy_kwh > aircraft.usable_energy_kwh * (1 - aircraft.reserve_fraction):
        return 0.0, best
    for _ in range(60):
        middle = (low + high) / 2
        flight = flight_for_payload(aircraft, energy_per_meter, base, site, maximum_ground_elevation, middle)
        one_way_distance = flight.horizontal_distance_m / 2
        range_feasible = (
            one_way_distance <= flight.range_budget_m
            and one_way_distance <= aircraft.empty_range_m
        )
        energy_feasible = flight.energy_kwh <= aircraft.usable_energy_kwh * (1 - aircraft.reserve_fraction)
        if range_feasible and energy_feasible:
            low, best = middle, flight
        else:
            high = middle
    return low, best


def _is_feasible(
    boxes: tuple[Box, ...], aircraft: Aircraft, safe_payload: float, flight: Flight
) -> bool:
    mass = sum(box.mass_kg for box in boxes)
    volume = sum(box.volume_m3 for box in boxes)
    energy_limit = aircraft.usable_energy_kwh * (1 - aircraft.reserve_fraction)
    return (
        mass <= min(aircraft.max_payload_kg, safe_payload) + 1e-9
        and volume <= aircraft.volume_m3 + 1e-12
        and flight.energy_kwh <= energy_limit + 1e-12
        and flight.horizontal_distance_m / 2 <= flight.range_budget_m + 1e-9
        and flight.horizontal_distance_m / 2 <= aircraft.empty_range_m + 1e-9
    )


def _validate_solution(boxes: list[Box], chosen: list[Candidate]) -> None:
    expected = {box.code for box in boxes}
    assigned = [code for candidate in chosen for code in candidate.boxes]
    if len(assigned) != len(set(assigned)) or set(assigned) != expected:
        raise AssertionError("Set-partition solution must assign every box exactly once")


def build_candidates(
    site_code: str,
    boxes: list[Box],
    base: Node,
    site: Node,
    dem_max: float,
    aircrafts: dict[str, Aircraft],
    energy_rates: dict[str, float],
    reserve_overrides: dict[str, float] | None = None,
) -> list[Candidate]:
    candidates = []
    aircraft_list = []
    for original in aircrafts.values():
        if reserve_overrides and original.code in reserve_overrides:
            aircraft = Aircraft(**{**asdict(original), "reserve_fraction": reserve_overrides[original.code]})
        else:
            aircraft = original
        aircraft_list.append(aircraft)
    safe_payload_by_aircraft = {
        aircraft.code: max_safe_payload(
            aircraft, energy_rates[aircraft.code], base, site, dem_max
        )[0]
        for aircraft in aircraft_list
    }
    for mask in range(1, 1 << len(boxes)):
        subset = tuple(boxes[index] for index in range(len(boxes)) if mask & (1 << index))
        mass = sum(box.mass_kg for box in subset)
        volume = sum(box.volume_m3 for box in subset)
        for aircraft in aircraft_list:
            if mass > aircraft.max_payload_kg or volume > aircraft.volume_m3:
                continue
            flight = flight_for_payload(
                aircraft, energy_rates[aircraft.code], base, site, dem_max, mass
            )
            safe_payload = safe_payload_by_aircraft[aircraft.code]
            if not _is_feasible(subset, aircraft, safe_payload, flight):
                continue
            work_time = (
                aircraft.prep_s + aircraft.load_each_s * len(subset)
                + flight.time_s + aircraft.handoff_base_s
                + aircraft.handoff_each_s * len(subset)
            )
            candidates.append(Candidate(
                site=site_code,
                aircraft=aircraft.code,
                boxes=tuple(box.code for box in subset),
                mass_kg=mass,
                volume_m3=volume,
                energy_kwh=flight.energy_kwh,
                outbound_energy_kwh=flight.outbound_energy_kwh,
                return_energy_kwh=flight.return_energy_kwh,
                flight_time_s=flight.time_s,
                outbound_time_s=flight.outbound_time_s,
                return_time_s=flight.return_time_s,
                work_time_s=work_time,
                range_budget_m=flight.range_budget_m,
                return_range_budget_m=flight.return_range_budget_m,
                reserve_kwh=aircraft.usable_energy_kwh * (1-aircraft.reserve_fraction)-flight.energy_kwh,
                return_soc_percent=100 * (1 - flight.energy_kwh / aircraft.usable_energy_kwh),
                feasible=True,
            ))
    return candidates


def solve_set_partition(
    boxes: list[Box],
    candidates: list[Candidate],
    objective_order: tuple[int, int, int] = (0, 1, 2),
) -> list[Candidate]:
    """Solve an exact set partition with a selected lexicographic objective order."""
    count = len(boxes)
    full_mask = (1 << count) - 1
    index = {box.code: position for position, box in enumerate(boxes)}
    by_first: list[list[tuple[int, Candidate]]] = [[] for _ in boxes]
    for candidate in candidates:
        mask = 0
        for code in candidate.boxes:
            mask |= 1 << index[code]
        first_index = min(index[code] for code in candidate.boxes)
        by_first[first_index].append((mask, candidate))

    best: dict[int, tuple[tuple[int, float, float], tuple[Candidate, ...]]] = {0: ((0, 0.0, 0.0), ())}
    for mask in range(full_mask + 1):
        if mask not in best:
            continue
        score, chosen = best[mask]
        remaining = full_mask ^ mask
        if remaining == 0:
            continue
        first_bit = (remaining & -remaining).bit_length() - 1
        for candidate_mask, candidate in by_first[first_bit]:
            if candidate_mask & mask:
                continue
            new_mask = mask | candidate_mask
            values = (1, candidate.energy_kwh, candidate.work_time_s)
            new_score = tuple(score[index] + values[index] for index in range(3))
            previous = best.get(new_mask)
            if previous is None or tuple(new_score[index] for index in objective_order) < tuple(previous[0][index] for index in objective_order):
                best[new_mask] = (new_score, chosen + (candidate,))
    if full_mask not in best:
        raise ValueError(f"No feasible complete partition for {boxes[0].site}")
    return list(best[full_mask][1])


def route_max_elevation(dem: rasterio.io.DatasetReader, base: Node, site: Node) -> tuple[float, float]:
    maximum, distance = sample_leg(dem, base, site)
    return maximum, distance


def run(
    reserves: Iterable[float] | None = None,
    energy_scales: Iterable[float] | None = None,
) -> dict:
    aircrafts = load_aircraft()
    base, sites = load_nodes()
    boxes_by_site = load_boxes()
    default_reserves = {aircraft.reserve_fraction for aircraft in aircrafts.values()}
    if len(default_reserves) != 1:
        raise ValueError("Aircraft workbook reserve defaults differ; define per-type baseline explicitly")
    default_reserve = default_reserves.pop()
    reserve_levels = [default_reserve]
    if reserves is not None:
        reserve_levels.extend(float(reserve) for reserve in reserves)
    reserve_levels = list(dict.fromkeys(round(value, 10) for value in reserve_levels))
    energy_scale_levels = [1.0]
    if energy_scales is not None:
        energy_scale_levels.extend(float(scale) for scale in energy_scales)
    energy_scale_levels = list(dict.fromkeys(round(value, 10) for value in energy_scale_levels))
    if any(scale <= 0 for scale in energy_scale_levels):
        raise ValueError("Horizontal energy scale factors must be positive")
    base_energy_rates = load_energy_per_meter(aircrafts)
    scenarios = [(default_reserve, 1.0)]
    scenarios.extend((reserve, 1.0) for reserve in reserve_levels if reserve != default_reserve)
    scenarios.extend((default_reserve, scale) for scale in energy_scale_levels if scale != 1.0)
    all_runs = []
    with rasterio.open(DEM_PATH) as dem:
        for reserve, energy_scale in scenarios:
            energy_rates = {code: rate * energy_scale for code, rate in base_energy_rates.items()}
            reserve_overrides = {code: reserve for code in aircrafts}
            payload_rows = []
            all_candidates: dict[str, list[Candidate]] = {}
            route_metadata = {}
            for site_code, site in sites.items():
                max_elevation, distance = route_max_elevation(dem, base, site)
                route_metadata[site_code] = {"distance_m": distance, "max_ground_elevation_m": max_elevation}
                for aircraft in aircrafts.values():
                    active = Aircraft(**{**asdict(aircraft), "reserve_fraction": reserve})
                    safe_payload, flight = max_safe_payload(
                        active, energy_rates[active.code], base, site, max_elevation
                    )
                    payload_rows.append({
                        "site": site_code, "aircraft": active.code,
                        "reserve_fraction": reserve,
                        "max_safe_payload_kg": safe_payload,
                        "route_distance_m": distance,
                        "cruise_altitude_m": flight.cruise_altitude_m,
                        "empty_round_trip_energy_kwh": flight.energy_kwh,
                    })
                candidates = build_candidates(
                    site_code, boxes_by_site[site_code], base, site, max_elevation,
                    aircrafts, energy_rates, reserve_overrides,
                )
                all_candidates[site_code] = candidates
            objective_policies = {
                "架次优先": (0, 1, 2),
                "能耗优先": (1, 0, 2),
                "时间优先": (2, 0, 1),
            }
            policy_partitions = {}
            infeasible_sites = set()
            for name, order in objective_policies.items():
                partitions = {}
                for site_code in sites:
                    try:
                        partitions[site_code] = solve_set_partition(
                            boxes_by_site[site_code], all_candidates[site_code], order
                        )
                    except ValueError:
                        infeasible_sites.add(site_code)
                        break
                if len(partitions) == len(sites):
                    policy_partitions[name] = partitions
            if infeasible_sites:
                policy_partitions = {}
            for partitions in policy_partitions.values():
                for site_code, chosen in partitions.items():
                    _validate_solution(boxes_by_site[site_code], chosen)
            primary = policy_partitions["架次优先"]
            solution_rows = []
            for policy, partitions in policy_partitions.items():
                chosen_candidates = [candidate for values in partitions.values() for candidate in values]
                solution_rows.append({
                    "policy": policy,
                    "flight_count": len(chosen_candidates),
                    "total_energy_kwh": sum(candidate.energy_kwh for candidate in chosen_candidates),
                    "total_work_time_s": sum(candidate.work_time_s for candidate in chosen_candidates),
                    "partitions": partitions,
                })
            nondominated = []
            for candidate_solution in solution_rows:
                score = (candidate_solution["flight_count"], candidate_solution["total_energy_kwh"], candidate_solution["total_work_time_s"])
                dominated = any(
                    other is not candidate_solution
                    and all(a <= b + 1e-9 for a, b in zip(
                        (other["flight_count"], other["total_energy_kwh"], other["total_work_time_s"]), score
                    ))
                    and any(a < b - 1e-9 for a, b in zip(
                        (other["flight_count"], other["total_energy_kwh"], other["total_work_time_s"]), score
                    ))
                    for other in solution_rows
                )
                if not dominated:
                    nondominated.append(candidate_solution["policy"])
            primary_flat = [candidate for candidates in primary.values() for candidate in candidates]
            all_runs.append({
                "reserve_fraction": reserve,
                "reserve_source": "aircraft_workbook_default" if reserve == default_reserve else "sensitivity_override",
                "horizontal_energy_scale": energy_scale,
                "complete_delivery_feasible": bool(policy_partitions),
                "infeasible_sites": sorted(infeasible_sites),
                "safe_payloads": payload_rows,
                "route_metadata": route_metadata,
                "partitions": {
                    site: [asdict(candidate) for candidate in candidates]
                    for site, candidates in primary.items()
                },
                "objective_solutions": [
                    {key: value for key, value in solution.items() if key != "partitions"}
                    for solution in solution_rows
                ],
                "sampled_nondominated_policies": nondominated,
                "summary": {
                    "box_count": sum(len(items) for items in boxes_by_site.values()),
                    "flight_count": len(primary_flat) if policy_partitions else None,
                    "total_energy_kwh": sum(candidate.energy_kwh for candidate in primary_flat) if policy_partitions else None,
                    "total_work_time_s": sum(candidate.work_time_s for candidate in primary_flat) if policy_partitions else None,
                },
            })
    return {
        "model": {
            "energy_calibration": "usable_energy_kwh / empty_standard_range_m as baseline horizontal energy per meter; check one-way range against payload-adjusted outbound and empty return ranges",
            "horizontal_energy_sensitivity": "horizontal calibrated energy rate multiplied by the scenario scale; 1.0 is baseline",
            "reserve_basis": "baseline uses the 20% return SOC lower bound in the aircraft workbook; sensitivity scenarios apply a common fraction to all types",
            "objective_tradeoff": "three lexicographic priority policies are compared; nondominated among these representative solutions, not a proof of the complete Pareto frontier",
            "dem": str(DEM_PATH.relative_to(ROOT)),
            "crs": "EPSG:4326; great-circle distance used for horizontal legs",
        },
        "runs": all_runs,
    }


def write_outputs(result: dict, output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    (output_dir / "problem1_results.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    run_data = result["runs"][0]
    with (output_dir / "safe_payloads.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=run_data["safe_payloads"][0].keys(), lineterminator="\n")
        writer.writeheader()
        writer.writerows(run_data["safe_payloads"])
    chosen_rows = [candidate for batches in run_data["partitions"].values() for candidate in batches]
    template_fields = [
        "架次编号", "服务区编号", "机型编号", "货箱编号列表", "总质量（kg）",
        "总体积（m³）", "往返时间（s）", "架次能耗（kWh）", "返航SOC（%）",
    ]
    with (output_dir / "batches.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=template_fields, lineterminator="\n")
        writer.writeheader()
        for index, row in enumerate(chosen_rows, start=1):
            writer.writerow({
                "架次编号": f"Q1-{index:03d}",
                "服务区编号": row["site"],
                "机型编号": row["aircraft"],
                "货箱编号列表": ",".join(row["boxes"]),
                "总质量（kg）": row["mass_kg"],
                "总体积（m³）": row["volume_m3"],
                "往返时间（s）": row["flight_time_s"],
                "架次能耗（kWh）": row["energy_kwh"],
                "返航SOC（%）": row["return_soc_percent"],
            })

    detail_fields = [
        "架次编号", "服务区编号", "机型编号", "货箱编号列表", "总质量（kg）", "总体积（m³）",
        "去程能耗（kWh）", "返程能耗（kWh）", "往返能耗（kWh）", "去程时间（s）",
        "返程时间（s）", "作业时间（s）", "返航SOC（%）", "允许返航SOC（%）", "可行状态",
    ]
    aircraft_defaults = {item["aircraft"]: item["reserve_fraction"] for item in run_data["safe_payloads"]}
    with (output_dir / "batches_detailed.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=detail_fields, lineterminator="\n")
        writer.writeheader()
        for index, row in enumerate(chosen_rows, start=1):
            writer.writerow({
                "架次编号": f"Q1-{index:03d}", "服务区编号": row["site"], "机型编号": row["aircraft"],
                "货箱编号列表": ",".join(row["boxes"]), "总质量（kg）": row["mass_kg"],
                "总体积（m³）": row["volume_m3"], "去程能耗（kWh）": row["outbound_energy_kwh"],
                "返程能耗（kWh）": row["return_energy_kwh"], "往返能耗（kWh）": row["energy_kwh"],
                "去程时间（s）": row["outbound_time_s"], "返程时间（s）": row["return_time_s"],
                "作业时间（s）": row["work_time_s"], "返航SOC（%）": row["return_soc_percent"],
                "允许返航SOC（%）": aircraft_defaults[row["aircraft"]] * 100,
                "可行状态": "可行" if row["feasible"] else "不可行",
            })

    with (output_dir / "service_area_summary.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        fields = ["服务区编号", "货箱数", "往返架次数", "总能耗（kWh）", "累计作业时间（s）", "未交付箱数"]
        writer = csv.DictWriter(stream, fieldnames=fields, lineterminator="\n")
        writer.writeheader()
        for site, batches in run_data["partitions"].items():
            writer.writerow({
                "服务区编号": site,
                "货箱数": sum(len(batch["boxes"]) for batch in batches),
                "往返架次数": len(batches),
                "总能耗（kWh）": sum(batch["energy_kwh"] for batch in batches),
                "累计作业时间（s）": sum(batch["work_time_s"] for batch in batches),
                "未交付箱数": 0,
            })

    with (output_dir / "objective_tradeoff.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        fields = ["返航余量（%）", "水平能耗率倍率", "目标优先策略", "总架次", "总能耗（kWh）", "累计作业时间（s）", "代表方案集中非支配"]
        writer = csv.DictWriter(stream, fieldnames=fields, lineterminator="\n")
        writer.writeheader()
        for item in result["runs"]:
            for solution in item["objective_solutions"]:
                writer.writerow({
                    "返航余量（%）": item["reserve_fraction"] * 100,
                    "水平能耗率倍率": item["horizontal_energy_scale"],
                    "目标优先策略": solution["policy"],
                    "总架次": solution["flight_count"],
                    "总能耗（kWh）": solution["total_energy_kwh"],
                    "累计作业时间（s）": solution["total_work_time_s"],
                    "代表方案集中非支配": "是" if solution["policy"] in item["sampled_nondominated_policies"] else "否",
                })
    with (output_dir / "sensitivity_summary.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        fields = ["reserve_fraction", "horizontal_energy_scale", "feasible", "infeasible_sites", "box_count", "flight_count", "total_energy_kwh", "total_work_time_s"]
        writer = csv.DictWriter(stream, fieldnames=fields, lineterminator="\n")
        writer.writeheader()
        for item in result["runs"]:
            writer.writerow({
                "reserve_fraction": item["reserve_fraction"],
                "horizontal_energy_scale": item["horizontal_energy_scale"],
                "feasible": item["complete_delivery_feasible"],
                "infeasible_sites": ",".join(item["infeasible_sites"]),
                **item["summary"],
            })

    with (output_dir / "safe_payloads_sensitivity.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        fields = ["reserve_fraction", "horizontal_energy_scale", "site", "aircraft", "max_safe_payload_kg", "route_distance_m", "cruise_altitude_m", "empty_round_trip_energy_kwh"]
        writer = csv.DictWriter(stream, fieldnames=fields, lineterminator="\n")
        writer.writeheader()
        for item in result["runs"]:
            for payload in item["safe_payloads"]:
                writer.writerow({"reserve_fraction": item["reserve_fraction"], "horizontal_energy_scale": item["horizontal_energy_scale"], **payload})

    plt.rcParams["font.family"] = "Noto Sans CJK SC"
    plt.rcParams["axes.unicode_minus"] = False
    _plot_safe_payloads(run_data, output_dir / "safe_payloads_heatmap.png")
    _plot_sensitivity(result["runs"], output_dir / "reserve_sensitivity.png")
    _plot_energy_sensitivity(result["runs"], output_dir / "energy_model_sensitivity.png")
    _plot_site_flights(run_data, output_dir / "site_flight_counts.png")
    _write_submission_workbook(chosen_rows, output_dir / "problem1_submission.xlsx")


def _write_submission_workbook(chosen_rows: list[dict], output_path: Path) -> None:
    template_path = ROOT / "docs" / "结果提交模板.xlsx"
    workbook = load_workbook(template_path)
    worksheet = workbook["Q1_单点组批"]
    headers = [cell.value for cell in worksheet[1]]
    worksheet.delete_rows(2, max(worksheet.max_row - 1, 1))
    for index, batch in enumerate(chosen_rows, start=1):
        worksheet.append([
            f"Q1-{index:03d}", batch["site"], batch["aircraft"],
            ",".join(batch["boxes"]), batch["mass_kg"], batch["volume_m3"],
            batch["flight_time_s"], batch["energy_kwh"], batch["return_soc_percent"],
        ])
    actual_headers = [cell.value for cell in worksheet[1]]
    if actual_headers[:9] != headers[:9]:
        raise AssertionError("Problem 1 submission workbook headers changed unexpectedly")
    workbook.save(output_path)


def _plot_safe_payloads(run_data: dict, output_path: Path) -> None:
    rows = run_data["safe_payloads"]
    sites = list(dict.fromkeys(row["site"] for row in rows))
    aircraft_types = ["A", "B", "C"]
    values = np.array([
        [next(row["max_safe_payload_kg"] for row in rows if row["site"] == site and row["aircraft"] == aircraft)
         for site in sites]
        for aircraft in aircraft_types
    ])
    figure, axis = plt.subplots(figsize=(14, 4.2), layout="constrained")
    image = axis.imshow(values, cmap="YlGnBu", aspect="auto", vmin=0)
    axis.set_xticks(range(len(sites)), sites)
    axis.set_yticks(range(len(aircraft_types)), [f"机型 {code}" for code in aircraft_types])
    axis.set_title("各服务区—机型最大安全载荷（kg）")
    axis.set_xlabel("服务区")
    for row_index in range(values.shape[0]):
        for column_index in range(values.shape[1]):
            axis.text(column_index, row_index, f"{values[row_index, column_index]:.1f}",
                      ha="center", va="center", fontsize=8,
                      color="white" if values[row_index, column_index] > values.max() * 0.58 else "black")
    figure.colorbar(image, ax=axis, label="kg")
    figure.savefig(output_path, dpi=180)
    plt.close(figure)


def _plot_sensitivity(runs: list[dict], output_path: Path) -> None:
    runs = sorted(
        (item for item in runs if math.isclose(item["horizontal_energy_scale"], 1.0)),
        key=lambda item: item["reserve_fraction"],
    )
    if not runs:
        return
    reserve = [item["reserve_fraction"] * 100 for item in runs]
    flights = [item["summary"]["flight_count"] for item in runs]
    energies = [item["summary"]["total_energy_kwh"] for item in runs]
    times = [item["summary"]["total_work_time_s"] / 3600 for item in runs]
    figure, axes = plt.subplots(1, 2, figsize=(10, 4), layout="constrained")
    axes[0].plot(reserve, flights, marker="o", color="#1769aa", linewidth=2)
    axes[0].set(title="返航余量与总架次数", xlabel="返航安全余量（%）", ylabel="往返架次")
    axes[0].grid(alpha=0.25)
    axes[1].plot(reserve, energies, marker="o", label="总能耗（kWh）", color="#c0392b")
    axes[1].set(xlabel="返航安全余量（%）", ylabel="总能耗（kWh）")
    second_axis = axes[1].twinx()
    second_axis.plot(reserve, times, marker="s", linestyle="--", label="累计作业时间（h）", color="#16825d")
    second_axis.set_ylabel("累计作业时间（h）")
    axes[1].set_title("返航余量与能耗、累计作业时间")
    axes[1].grid(alpha=0.25)
    handles = axes[1].get_lines() + second_axis.get_lines()
    axes[1].legend(handles, [line.get_label() for line in handles], loc="best", fontsize=8)
    figure.savefig(output_path, dpi=180)
    plt.close(figure)


def _plot_energy_sensitivity(runs: list[dict], output_path: Path) -> None:
    runs = sorted(
        (item for item in runs if math.isclose(item["reserve_fraction"], 0.2)),
        key=lambda item: item["horizontal_energy_scale"],
    )
    if not runs:
        return
    scales = [item["horizontal_energy_scale"] for item in runs]
    payload_values = [
        min(row["max_safe_payload_kg"] for row in item["safe_payloads"])
        for item in runs
    ]
    flight_counts = [item["summary"]["flight_count"] for item in runs]
    figure, axes = plt.subplots(1, 2, figsize=(10, 4), layout="constrained")
    axes[0].plot(scales, payload_values, marker="o", color="#1769aa")
    axes[0].set(title="能耗假设倍率与最小安全载荷", xlabel="水平能耗率倍率（基准=1）", ylabel="15×3组合中的最小安全载荷（kg）")
    axes[0].grid(alpha=0.25)
    axes[1].plot(scales, flight_counts, marker="s", color="#c0392b")
    axes[1].set(title="能耗假设倍率与组批架次", xlabel="水平能耗率倍率（基准=1）", ylabel="往返架次")
    axes[1].grid(alpha=0.25)
    figure.savefig(output_path, dpi=180)
    plt.close(figure)


def _plot_site_flights(run_data: dict, output_path: Path) -> None:
    sites = list(run_data["partitions"])
    counts = [len(run_data["partitions"][site]) for site in sites]
    figure, axis = plt.subplots(figsize=(11, 4), layout="constrained")
    bars = axis.bar(sites, counts, color="#3478a8")
    axis.bar_label(bars, padding=2)
    axis.set(title="基准返航余量下各服务区往返架次数", xlabel="服务区", ylabel="架次")
    axis.grid(axis="y", alpha=0.25)
    figure.savefig(output_path, dpi=180)
    plt.close(figure)


def main() -> None:
    parser = argparse.ArgumentParser(description="Solve Problem 1 single-site round-trip batches")
    parser.add_argument("--reserve", nargs="+", type=float, help="Override each type's reserve fraction, e.g. 0.1 0.2 0.3")
    parser.add_argument("--energy-scale", nargs="+", type=float, help="Scale the assumed horizontal energy rate, e.g. 0.8 1.2")
    parser.add_argument("--output", type=Path, default=ROOT / "outputs/problem1")
    arguments = parser.parse_args()
    result = run(arguments.reserve, arguments.energy_scale)
    write_outputs(result, arguments.output)
    for item in result["runs"]:
        print(item["reserve_fraction"], item["horizontal_energy_scale"], item["complete_delivery_feasible"], item["summary"])
    print(f"Wrote results to {arguments.output}")


if __name__ == "__main__":
    main()
