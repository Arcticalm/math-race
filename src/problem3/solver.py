from __future__ import annotations

import argparse
import copy
import csv
import json
import math
from dataclasses import asdict, dataclass
from pathlib import Path

import rasterio
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import load_workbook

from src.problem1.solver import Node, load_aircraft, sample_leg
from src.problem2.solver import (
    RouteEvaluator,
    _hard_deadlines,
    _validate,
    _charge_duration as transport_charge_duration,
    load_resources,
    solve as solve_problem2,
)
from src.problem3.physics import (
    LinkEvaluator,
    sampled_flight_leg,
    estimate_relay_mission,
    link_limits,
    load_link_parameters,
    load_relay_parameters,
)

ROOT = Path(__file__).resolve().parents[2]
OUTPUT_DEFAULT = ROOT / "outputs" / "problem3"


@dataclass(frozen=True)
class TrackPhase:
    sortie: str
    name: str
    start_s: float
    end_s: float
    start_node: Node
    end_node: Node
    start_altitude_m: float
    end_altitude_m: float


def _node_at(first: Node, second: Node, fraction: float, code: str) -> Node:
    return Node(code, first.longitude + fraction * (second.longitude - first.longitude),
                first.latitude + fraction * (second.latitude - first.latitude), 0.0)


def build_trajectory(sorties, reserve: float = 0.2, energy_scale: float = 1.0) -> list[TrackPhase]:
    aircrafts = load_aircraft()
    base_evaluator = RouteEvaluator(reserve, energy_scale)
    try:
        base, sites = base_evaluator.base, base_evaluator.sites
        phases = []
        for sortie in sorties:
            aircraft = aircrafts[sortie.aircraft]
            current_node = base
            current_altitude = base.elevation_m
            current_time = sortie.takeoff_s
            for sequence, leg in enumerate(sortie.legs or [], 1):
                destination = base if leg["to"] == "O01" else sites[leg["to"]]
                destination_altitude = destination.elevation_m + (30 if destination.code != "O01" else 0)
                cruise_altitude = leg["cruise_altitude_m"]
                distance = leg["distance_m"]
                climb_s = max(0.0, cruise_altitude - current_altitude) / aircraft.climb_speed_mps
                cruise_s = distance / aircraft.cruise_speed_mps
                descent_s = max(0.0, cruise_altitude - destination_altitude) / aircraft.descent_speed_mps
                phases.append(TrackPhase(sortie.code, f"leg{sequence}-climb", current_time,
                    current_time + climb_s, current_node, current_node, current_altitude, cruise_altitude))
                current_time += climb_s
                phases.append(TrackPhase(sortie.code, f"leg{sequence}-cruise", current_time,
                    current_time + cruise_s, current_node, destination, cruise_altitude, cruise_altitude))
                current_time += cruise_s
                phases.append(TrackPhase(sortie.code, f"leg{sequence}-descent", current_time,
                    current_time + descent_s, destination, destination, cruise_altitude, destination_altitude))
                current_time += descent_s
                current_node, current_altitude = destination, destination_altitude
                if destination.code != "O01":
                    count = sum(box.site == destination.code for box in sortie.boxes)
                    handoff = aircraft.handoff_base_s + aircraft.handoff_each_s * count
                    phases.append(TrackPhase(sortie.code, f"{destination.code}-handoff", current_time,
                        current_time + handoff, destination, destination, destination_altitude, destination_altitude))
                    current_time += handoff
            if abs(current_time - sortie.return_s) > 1e-6:
                raise ValueError(f"Trajectory timeline mismatch for {sortie.code}: {current_time} != {sortie.return_s}")
        return phases
    finally:
        base_evaluator.close()


def _position(phase: TrackPhase, time_s: float) -> tuple[Node, float]:
    duration = phase.end_s - phase.start_s
    fraction = 0.0 if duration <= 0 else min(1.0, max(0.0, (time_s - phase.start_s) / duration))
    node = _node_at(phase.start_node, phase.end_node, fraction, phase.name)
    altitude = phase.start_altitude_m + fraction * (phase.end_altitude_m - phase.start_altitude_m)
    return node, altitude


def screen_direct_links(sorties, phases: list[TrackPhase], step_s: float = 10.0) -> tuple[list[dict], list[dict]]:
    params = load_link_parameters()
    evaluator = LinkEvaluator(params)
    base, _ = evaluator.base, evaluator.sites
    gateway = Node("G01", base.longitude, base.latitude, base.elevation_m + 20)
    threshold = link_limits(params)["transport_gateway_db"]
    samples = []
    try:
        for phase in phases:
            if phase.end_s <= phase.start_s:
                continue
            sample_count = max(1, math.ceil((phase.end_s - phase.start_s) / step_s))
            for index in range(sample_count + 1):
                time_s = phase.start_s + (phase.end_s - phase.start_s) * index / sample_count
                if index == sample_count and phase.end_s == phase.start_s:
                    continue
                position, altitude = _position(phase, time_s)
                link = evaluator.evaluate(position, altitude, gateway, gateway.elevation_m, threshold)
                samples.append({
                    "sortie": phase.sortie, "phase": phase.name, "time_s": time_s,
                    "longitude": position.longitude, "latitude": position.latitude,
                    "altitude_m": altitude, **asdict(link),
                })
    finally:
        evaluator.close()
    intervals = _sample_intervals(samples)
    return samples, intervals


def _sample_intervals(samples: list[dict]) -> list[dict]:
    """Cover every sample cell; either failed endpoint requests relay coverage."""
    groups = {}
    for sample in samples:
        groups.setdefault((sample["sortie"], sample["phase"]), []).append(sample)
    intervals = []
    for (sortie, phase), points in groups.items():
        points.sort(key=lambda item: item["time_s"])
        for first, second in zip(points, points[1:]):
            if second["time_s"] <= first["time_s"]:
                continue
            available = first["available"] and second["available"]
            reason = first["reason"] or second["reason"]
            if (intervals and intervals[-1]["sortie"] == sortie
                    and intervals[-1]["phase"] == phase
                    and intervals[-1]["direct_available"] == available
                    and intervals[-1]["reason"] == reason):
                intervals[-1]["end_s"] = second["time_s"]
            else:
                intervals.append({"sortie": sortie, "phase": phase,
                                  "start_s": first["time_s"], "end_s": second["time_s"],
                                  "direct_available": available, "reason": reason})
    return sorted(intervals, key=lambda item: (item["start_s"], item["sortie"]))


def _peak_gap_load(intervals: list[dict]) -> tuple[int, float]:
    events = []
    for item in intervals:
        if item["direct_available"] or item["end_s"] <= item["start_s"]:
            continue
        events.append((item["start_s"], 1))
        events.append((item["end_s"], -1))
    events.sort(key=lambda item: (item[0], item[1]))
    active = peak = 0
    previous = None
    squared_load = 0.0
    for time_s, change in events:
        if previous is not None and time_s > previous:
            squared_load += (time_s - previous) * active * active
        active += change
        peak = max(peak, active)
        previous = time_s
    return peak, squared_load


def _merge_gap_intervals(intervals: list[dict]) -> list[dict]:
    merged = []
    by_sortie = {}
    for interval in intervals:
        if interval["direct_available"] or interval["end_s"] <= interval["start_s"]:
            continue
        by_sortie.setdefault(interval["sortie"], []).append(interval)
    for sortie, values in by_sortie.items():
        values.sort(key=lambda item: item["start_s"])
        current = None
        for value in values:
            if current is None or value["start_s"] > current["end_s"] + 1e-6:
                if current is not None:
                    merged.append(current)
                current = {
                    "sortie": value["sortie"], "start_s": value["start_s"],
                    "end_s": value["end_s"], "direct_available": False,
                    "reason": value.get("reason", ""),
                }
            else:
                current["end_s"] = max(current["end_s"], value["end_s"])
        if current is not None:
            merged.append(current)
    return sorted(merged, key=lambda item: (item["start_s"], item["sortie"]))


def _assign_gap_ids(intervals: list[dict], gaps: list[dict]) -> None:
    for gap_id, gap in enumerate(gaps, 1):
        gap["gap_id"] = gap_id
    for interval in intervals:
        interval.pop("gap_id", None)
        if interval["direct_available"]:
            continue
        matches = [gap["gap_id"] for gap in gaps
                   if gap["sortie"] == interval["sortie"]
                   and interval["start_s"] < gap["end_s"] - 1e-7
                   and interval["end_s"] > gap["start_s"] + 1e-7]
        if matches:
            interval["gap_id"] = matches[0]


def _transport_schedule_resource_feasible(sorties, batteries) -> bool:
    by_drone: dict[str, list[tuple[float, float]]] = {}
    by_battery: dict[str, list[tuple[float, float]]] = {}
    battery_by_id = {item.code: item for item in batteries}
    aircrafts = load_aircraft()
    for sortie in sorties:
        by_drone.setdefault(sortie.drone, []).append((sortie.prep_start_s, sortie.return_s))
        battery = battery_by_id[sortie.battery]
        soc = 1 - sortie.energy_kwh / aircrafts[sortie.aircraft].usable_energy_kwh
        available = sortie.return_s + transport_charge_duration(battery.full_charge_s, soc)
        by_battery.setdefault(sortie.battery, []).append((sortie.prep_start_s, available))
    for intervals in (*by_drone.values(), *by_battery.values()):
        intervals.sort()
        if any(previous[1] > current[0] + 1e-7 for previous, current in zip(intervals, intervals[1:])):
            return False
    return True


def _changed_sortie_resource_feasible(sortie, schedule, batteries) -> bool:
    battery_by_id = {item.code: item for item in batteries}
    aircrafts = load_aircraft()
    peers = [item for item in schedule if item.code != sortie.code]
    same_drone = sorted((item for item in peers if item.drone == sortie.drone),
                        key=lambda item: item.prep_start_s)
    if any(item.prep_start_s < sortie.return_s - 1e-7
           and item.return_s > sortie.prep_start_s + 1e-7 for item in same_drone):
        return False
    battery = battery_by_id[sortie.battery]
    soc = 1 - sortie.energy_kwh / aircrafts[sortie.aircraft].usable_energy_kwh
    available = sortie.return_s + transport_charge_duration(battery.full_charge_s, soc)
    same_battery = [item for item in peers if item.battery == sortie.battery]
    for item in same_battery:
        item_battery = battery_by_id[item.battery]
        item_soc = 1 - item.energy_kwh / aircrafts[item.aircraft].usable_energy_kwh
        item_available = item.return_s + transport_charge_duration(item_battery.full_charge_s, item_soc)
        if item.prep_start_s < available - 1e-7 and item_available > sortie.prep_start_s + 1e-7:
            return False
    return True


def _weighted_expected_lateness(sorties, boxes) -> float:
    delivered = {code: time_s for sortie in sorties for code, time_s in (sortie.deliveries or {}).items()}
    return sum(box.priority * max(0.0, delivered.get(box.code, 0.0) - box.due_s)
               for box in boxes if box.due_s is not None)


def _hard_deadline_audit(sorties) -> dict:
    checks = []
    for sortie in sorties:
        deliveries = sortie.deliveries or {}
        for box in sortie.boxes:
            delivered = deliveries.get(box.code)
            for kind, deadline_s in _hard_deadlines(box):
                checks.append({
                    "sortie": sortie.code, "box": box.code, "kind": kind,
                    "delivery_s": delivered, "deadline_s": deadline_s,
                    "passed": delivered is not None and delivered <= deadline_s + 1e-7,
                })
    return {
        "check_count": len(checks),
        "passed_count": sum(item["passed"] for item in checks),
        "failed_count": sum(not item["passed"] for item in checks),
        "passed": all(item["passed"] for item in checks),
        "checks": checks,
    }


def _shift_sortie(sortie, delay_s: float) -> None:
    sortie.prep_start_s += delay_s
    sortie.prep_end_s += delay_s
    sortie.loading_end_s += delay_s
    sortie.takeoff_s += delay_s
    sortie.return_s += delay_s
    sortie.deliveries = {code: time_s + delay_s for code, time_s in (sortie.deliveries or {}).items()}


def coordinate_transport_starts(sorties, boxes, intervals: list[dict], batteries,
                                step_s: float = 60.0, maximum_delay_s: float = 1800.0) -> tuple[list, dict]:
    schedule = copy.deepcopy(sorties)
    gaps_by_sortie: dict[str, list[dict]] = {}
    for interval in intervals:
        if not interval["direct_available"] and interval["end_s"] > interval["start_s"]:
            gaps_by_sortie.setdefault(interval["sortie"], []).append(dict(interval))
    max_delay = {}
    for sortie in schedule:
        slack = min((deadline_s - sortie.deliveries[box.code]
                     for box in sortie.boxes for _, deadline_s in _hard_deadlines(box)),
                    default=maximum_delay_s)
        if not math.isfinite(slack):
            slack = maximum_delay_s
        max_delay[sortie.code] = max(0.0, min(maximum_delay_s, slack))

    delays = {sortie.code: 0.0 for sortie in schedule}
    current_gaps = [item for values in gaps_by_sortie.values() for item in values]
    current_peak, current_load = _peak_gap_load(current_gaps)
    current_lateness = _weighted_expected_lateness(schedule, boxes)
    for _ in range(3):
        improved = False
        ordered = sorted(schedule, key=lambda sortie: (
            -sum(item["end_s"] - item["start_s"] for item in gaps_by_sortie.get(sortie.code, [])),
            sortie.prep_start_s,
        ))
        for sortie in ordered:
            original_delay = delays[sortie.code]
            best = (current_peak, current_load, current_lateness, sum(delays.values()), original_delay)
            best_delay = original_delay
            delay_values = [original_delay]
            delay_values.extend(value for value in range(0, int(max_delay[sortie.code]) + 1, int(step_s)))
            if max_delay[sortie.code] > 0:
                delay_values.append(max_delay[sortie.code])
            for candidate_delay in sorted(set(delay_values)):
                offset = candidate_delay - original_delay
                if abs(offset) < 1e-7:
                    continue
                candidate_schedule = copy.deepcopy(schedule)
                candidate_sortie = next(item for item in candidate_schedule if item.code == sortie.code)
                _shift_sortie(candidate_sortie, offset)
                if not _changed_sortie_resource_feasible(candidate_sortie, candidate_schedule, batteries):
                    continue
                candidate_gaps = []
                for code, values in gaps_by_sortie.items():
                    shift = candidate_delay if code == sortie.code else delays[code]
                    candidate_gaps.extend({**item, "start_s": item["start_s"] + shift,
                                           "end_s": item["end_s"] + shift} for item in values)
                peak, load = _peak_gap_load(candidate_gaps)
                lateness = _weighted_expected_lateness(candidate_schedule, boxes)
                score = (peak, load, lateness, sum(delays.values()) + offset, candidate_delay)
                if score < best:
                    best, best_delay = score, candidate_delay
            if best_delay != original_delay:
                offset = best_delay - original_delay
                _shift_sortie(sortie, offset)
                delays[sortie.code] = best_delay
                current_gaps = [
                    {**item, "start_s": item["start_s"] + delays[code],
                     "end_s": item["end_s"] + delays[code]}
                    for code, values in gaps_by_sortie.items() for item in values
                ]
                current_peak, current_load = _peak_gap_load(current_gaps)
                current_lateness = _weighted_expected_lateness(schedule, boxes)
                improved = True
        if not improved:
            break
    return schedule, {
        "method": "three-pass coordinate search over delayed sortie starts",
        "maximum_delay_s": maximum_delay_s,
        "step_s": step_s,
        "delays_s": {code: delay for code, delay in delays.items() if delay > 0},
        "peak_simultaneous_direct_gaps": current_peak,
        "integrated_squared_gap_load_s": current_load,
        "weighted_expected_lateness_s": current_lateness,
        "resource_schedule_feasible": _transport_schedule_resource_feasible(schedule, batteries),
    }


def search_relay_candidates(sorties, phases: list[TrackPhase], gaps: list[dict],
                            sample_step_s: float = 30.0,
                            candidate_window_s: float = 20000.0) -> list[dict]:
    params = load_link_parameters()
    relay = load_relay_parameters()
    evaluator = LinkEvaluator(params)
    base = evaluator.base
    gateway = Node("G01", base.longitude, base.latitude, base.elevation_m + 20)
    limits = link_limits(params)
    time_step_s = min(30.0, sample_step_s)
    offsets = (-0.01, 0.0, 0.01)
    results = []
    phases_by_sortie: dict[str, list[TrackPhase]] = {}
    sorties_by_code = {sortie.code: sortie for sortie in sorties}
    for phase in phases:
        phases_by_sortie.setdefault(phase.sortie, []).append(phase)
    try:
        gap_groups = []
        for gap_id, gap in enumerate(sorted(gaps, key=lambda item: (item["start_s"], item["sortie"])), 1):
            gap["gap_id"] = gap_id
            if (not gap_groups
                    or gap["start_s"] - gap_groups[-1][0]["start_s"] > candidate_window_s + 1e-7):
                gap_groups.append([gap])
            else:
                gap_groups[-1].append(gap)

        for group_index, group in enumerate(gap_groups, 1):
            demand_by_gap = {}
            target_nodes = []
            for gap in group:
                sortie_phases = phases_by_sortie[gap["sortie"]]
                demand_by_gap[gap["gap_id"]] = []
                sample_count = max(1, math.ceil((gap["end_s"] - gap["start_s"]) / time_step_s))
                for sample_index in range(sample_count + 1):
                    time_s = gap["start_s"] + (gap["end_s"] - gap["start_s"]) * sample_index / sample_count
                    phase = next((item for item in sortie_phases
                                  if item.start_s - 1e-7 <= time_s <= item.end_s + 1e-7), None)
                    if phase is None:
                        continue
                    transport_node, transport_altitude = _position(phase, time_s)
                    demand_by_gap[gap["gap_id"]].append((time_s, transport_node, transport_altitude))
                for fraction in (0.0, 0.25, 0.5, 0.75, 1.0):
                    time_s = gap["start_s"] + (gap["end_s"] - gap["start_s"]) * fraction
                    phase = next(item for item in sortie_phases
                                 if item.start_s - 1e-7 <= time_s <= item.end_s + 1e-7)
                    target_nodes.append(_position(phase, time_s)[0])
            locations = {(round(node.longitude + lon_offset, 7), round(node.latitude + lat_offset, 7))
                         for node in target_nodes for lon_offset in offsets for lat_offset in offsets}
            candidate_rows = []
            for longitude, latitude in sorted(locations):
                row, col = rasterio.transform.rowcol(evaluator.dem.transform, longitude, latitude)
                if not (0 <= row < evaluator.dem.height and 0 <= col < evaluator.dem.width):
                    continue
                ground_m = float(evaluator.dem_values[row, col])
                if not math.isfinite(ground_m) or math.isclose(ground_m, -32767.0):
                    continue
                hover = Node("H", longitude, latitude, ground_m)
                try:
                    outbound_max_ground, outbound_distance = sampled_flight_leg(evaluator, base, hover)
                    return_max_ground, return_distance = sampled_flight_leg(evaluator, hover, base)
                except ValueError:
                    continue
                for agl_m in (50.0, 150.0, 250.0, relay.maximum_agl_m):
                    if agl_m > relay.maximum_agl_m:
                        continue
                    hover = Node("H", longitude, latitude, ground_m)
                    hover_altitude = ground_m + agl_m
                    backhaul = evaluator.evaluate(hover, hover_altitude, gateway,
                                                   gateway.elevation_m, limits["relay_gateway_db"])
                    if not backhaul.available:
                        continue
                    covered = []
                    for gap in group:
                        access_ok = all(
                            evaluator.evaluate(transport_node, transport_altitude, hover,
                                               hover_altitude, limits["transport_relay_db"]).available
                            for _, transport_node, transport_altitude in demand_by_gap[gap["gap_id"]]
                        )
                        if access_ok:
                            covered.append(gap)
                    if not covered:
                        continue
                    for service_start in sorted({item["start_s"] for item in covered}):
                        service_subset = [item for item in covered
                                          if item["start_s"] >= service_start - 1e-7]
                        if not service_subset:
                            continue
                        for service_end in sorted({item["end_s"] for item in service_subset}):
                            window_subset = [item for item in service_subset
                                             if item["end_s"] <= service_end + 1e-7]
                            if not window_subset:
                                continue
                            if service_end <= service_start + 1e-7:
                                continue
                            service_s = service_end - service_start
                            mission = estimate_relay_mission(
                                service_s, outbound_distance, outbound_max_ground, hover_altitude,
                                return_distance, return_max_ground, base.elevation_m, relay,
                            )
                            prepare_start = service_start - relay.preparation_s - mission.outbound_flight_s - relay.link_setup_s
                            if prepare_start < -1e-7 or not mission.feasible_energy:
                                continue
                            subset_ids = [item["gap_id"] for item in window_subset]
                            candidate_rows.append({
                                "covered_gap_ids": subset_ids,
                                "covered_sorties": ",".join(sorted(item["sortie"] for item in window_subset)),
                                "service_start_s": service_start, "service_end_s": service_end,
                                "preparation_start_s": max(0.0, prepare_start),
                                "longitude": longitude, "latitude": latitude,
                                "ground_dsm_m": ground_m, "hover_altitude_m": hover_altitude,
                                "agl_m": agl_m, "energy_kwh": mission.energy_kwh,
                                "return_soc_percent": mission.return_soc_percent,
                                "return_o01_s": service_end + mission.return_flight_s,
                                "outbound_flight_s": mission.outbound_flight_s,
                                "return_flight_s": mission.return_flight_s,
                            })
            by_coverage = {}
            for candidate in candidate_rows:
                by_coverage.setdefault(tuple(candidate["covered_gap_ids"]), []).append(candidate)
            representatives = []
            for options in by_coverage.values():
                keys = (lambda c: c["energy_kwh"], lambda c: -c["preparation_start_s"],
                        lambda c: c["return_o01_s"])
                chosen = {}
                for key in keys:
                    for c in sorted(options, key=key)[:3]:
                        chosen[(c["longitude"], c["latitude"], c["hover_altitude_m"])] = c
                representatives.extend(chosen.values())
            candidate_rows = sorted(representatives, key=lambda item: (
                -len(item["covered_gap_ids"]), item["energy_kwh"], item["return_o01_s"],
            ))
            results.append({
                "group": group_index, "covered_sorties": ",".join(sorted({item["sortie"] for item in group})),
                "gap_count": len(group), "gap_start_s": min(item["start_s"] for item in group),
                "gap_end_s": max(item["end_s"] for item in group),
                "gap_start_by_id": {str(item["gap_id"]): item["start_s"] for item in group},
                "gaps": [{"gap_id": item["gap_id"], "sortie": item["sortie"],
                          "start_s": item["start_s"], "end_s": item["end_s"]} for item in group],
                "candidate_count": len(candidate_rows), "candidates": candidate_rows,
            })
    finally:
        evaluator.close()
    return results


def _write_csv(path: Path, rows: list[dict]) -> None:
    if not rows:
        path.write_text("", encoding="utf-8-sig")
        return
    fieldnames = list(dict.fromkeys(key for row in rows for key in row))
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fieldnames, extrasaction="ignore", lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def _relay_for_gap(selected: list[dict], gap_id: int) -> dict | None:
    for item in selected:
        if gap_id in item.get("covered_gap_ids", []):
            return item
    return None


def _communication_rows(sorties, intervals, selected) -> list[dict]:
    rows = []
    for interval in intervals:
        if interval["end_s"] <= interval["start_s"]:
            continue
        relay = None
        mode = "直连" if interval["direct_available"] else "中断"
        if not interval["direct_available"] and interval.get("gap_id") is not None:
            relay = _relay_for_gap(selected, int(interval["gap_id"]))
            if relay is not None:
                mode = "中继候选"
        rows.append({
            "运输架次编号": interval["sortie"],
            "通信阶段": interval.get("phase", "空中通信"),
            "开始时刻（s）": interval["start_s"],
            "结束时刻（s）": interval["end_s"],
            "保障方式": mode,
            "中继架次编号": relay.get("relay_sortie", "") if relay else "",
            "诊断说明": interval.get("reason", "") if mode == "中断" else "",
        })
    return rows


def _write_q3_submission(output_dir: Path, sorties, intervals, selected, communication_rows=None) -> None:
    workbook = load_workbook(ROOT / "docs" / "结果提交模板.xlsx")
    transport_sheet = workbook["Q2_运输架次"]
    if transport_sheet.max_row > 1:
        transport_sheet.delete_rows(2, transport_sheet.max_row - 1)
    for sortie in sorties:
        transport_sheet.append([
            sortie.code, sortie.drone, sortie.aircraft, sortie.battery,
            round(sortie.prep_start_s, 3), "→".join(["O01", *sortie.route, "O01"]),
            round(sortie.return_s, 3), round(sortie.energy_kwh, 6),
        ])
    delivery_sheet = workbook["Q2_逐箱交付"]
    if delivery_sheet.max_row > 1:
        delivery_sheet.delete_rows(2, delivery_sheet.max_row - 1)
    for sortie in sorties:
        for box in sortie.boxes:
            delivery_sheet.append([
                box.code, sortie.code, box.site,
                round((sortie.deliveries or {})[box.code], 3),
            ])
    relay_sheet = workbook["Q3_中继架次"]
    if relay_sheet.max_row > 1:
        relay_sheet.delete_rows(2, relay_sheet.max_row - 1)
    for item in selected:
        relay_sheet.append([
            item.get("relay_sortie", ""), item.get("relay_drone", ""),
            item.get("energy_component", ""), round(item["preparation_start_s"], 3),
            item["longitude"], item["latitude"], item["hover_altitude_m"],
            round(item["service_start_s"], 3), round(item["service_end_s"], 3),
            round(item["return_o01_s"], 3), round(item["energy_kwh"], 6),
        ])
    communication_sheet = workbook["Q3_通信保障"]
    if communication_sheet.max_row > 1:
        communication_sheet.delete_rows(2, communication_sheet.max_row - 1)
    for row in (communication_rows if communication_rows is not None else _communication_rows(sorties, intervals, selected)):
        communication_sheet.append([row[key] for key in (
            "运输架次编号", "通信阶段", "开始时刻（s）", "结束时刻（s）",
            "保障方式", "中继架次编号")])
    workbook.save(output_dir / "problem3_submission.xlsx")


def _plot_problem3_overview(output_path: Path, phases: list[TrackPhase], selected: list[dict]) -> None:
    figure, axis = plt.subplots(figsize=(10, 7), layout="constrained")
    seen = set()
    for phase in phases:
        if phase.sortie in seen or phase.start_node.code == phase.end_node.code:
            continue
        seen.add(phase.sortie)
        axis.plot([phase.start_node.longitude, phase.end_node.longitude],
                  [phase.start_node.latitude, phase.end_node.latitude],
                  linewidth=0.8, alpha=0.45)
    if selected:
        axis.scatter([item["longitude"] for item in selected],
                     [item["latitude"] for item in selected],
                     marker="^", s=55, color="#c0392b", label="Relay hover point")
    axis.set_xlabel("Longitude")
    axis.set_ylabel("Latitude")
    axis.set_title("Problem 3 transport routes and relay schedule")
    if selected:
        axis.legend()
    figure.savefig(output_path, dpi=160)
    plt.close(figure)


def _package_program(output_dir: Path) -> None:
    import zipfile
    files = [
        ROOT / "src/problem1/solver.py", ROOT / "src/problem1/__init__.py",
        ROOT / "src/problem2/solver.py", ROOT / "src/problem2/__init__.py",
        ROOT / "src/problem3/solver.py", ROOT / "src/problem3/physics.py",
        ROOT / "src/problem3/__init__.py", ROOT / "src/problem3/README.md",
        ROOT / "src/problem3/requirements.txt", ROOT / "src/problem3/audit.py", ROOT / "src/problem3/joint.py",
        ROOT / "src/problem1/requirements.txt", ROOT / "src/problem2/requirements.txt",
        ROOT / "tests/test_problem3.py", ROOT / "tests/__init__.py",
        ROOT / "docs/结果提交模板.xlsx",
    ]
    with zipfile.ZipFile(output_dir / "problem3_program.zip", "w", zipfile.ZIP_DEFLATED) as archive:
        for path in files:
            archive.write(path, path.relative_to(ROOT).as_posix())


def _node_elevation_audit(evaluator: LinkEvaluator) -> list[dict]:
    nodes = [evaluator.base, *evaluator.sites.values()]
    audit = []
    for node in nodes:
        row, col = rasterio.transform.rowcol(evaluator.dem.transform, node.longitude, node.latitude)
        valid = 0 <= row < evaluator.dem.height and 0 <= col < evaluator.dem.width
        dsm = float(evaluator.dem_values[row, col]) if valid else math.nan
        valid = valid and math.isfinite(dsm) and not math.isclose(dsm, -32767.0)
        audit.append({
            "node": node.code, "node_table_ground_m": node.elevation_m,
            "dsm_pixel_m": dsm if valid else "", "dsm_minus_node_m": dsm - node.elevation_m if valid else "",
            "valid": valid,
        })
    return audit


def _charge_duration(full_charge_s: float, soc: float) -> float:
    if soc < 0.9:
        return full_charge_s * (0.65 * (0.9 - soc) / 0.9 + 0.35)
    return full_charge_s * 0.35 * (1 - soc) / 0.1


def select_relay_schedule(candidate_groups: list[dict], gap_count: int,
                          search_limit: int = 50000, objective: str = "energy") -> dict:
    """Select interval missions with exact airframe and component capacities."""
    import numpy as np
    from scipy.optimize import Bounds, LinearConstraint, milp
    from scipy.sparse import lil_matrix

    relay = load_relay_parameters()
    candidates = [{**candidate, "candidate_group": group["group"]}
                  for group in candidate_groups for candidate in group["candidates"]]
    for candidate in candidates:
        if isinstance(candidate["covered_gap_ids"], str):
            candidate["covered_gap_ids"] = json.loads(candidate["covered_gap_ids"])
    impossible = [gap for gap in range(1, gap_count + 1)
                  if not any(gap in c["covered_gap_ids"] for c in candidates)]
    count = len(candidates)
    selected = []
    result = None
    if count:
        starts = np.array([c["preparation_start_s"] for c in candidates])
        drone_ends = np.array([c["return_o01_s"] + relay.turnaround_s for c in candidates])
        battery_ends = np.array([c["return_o01_s"] + _charge_duration(
            relay.full_charge_s, c["return_soc_percent"] / 100) for c in candidates])
        constraints = []
        for gap in range(1, gap_count + 1):
            if gap in impossible:
                continue
            constraints.append(([i for i, c in enumerate(candidates)
                                 if gap in c["covered_gap_ids"]], 1, np.inf))
        # An interval graph needs exactly its peak overlap number of colors.
        # Only starts change overlap upward; equal end/start releases first.
        for ends, capacity in ((drone_ends, 2), (battery_ends, relay.battery_count)):
            seen = set()
            for time_s in sorted(set(starts)):
                active = tuple(np.flatnonzero((starts <= time_s) & (ends > time_s + 1e-7)))
                if len(active) > capacity and active not in seen:
                    constraints.append((active, 0, capacity))
                    seen.add(active)
        matrix = lil_matrix((len(constraints), count))
        for row, (indices, _, _) in enumerate(constraints):
            matrix[row, list(indices)] = 1
        energies = np.array([c["energy_kwh"] for c in candidates])
        costs = energies + 1e-6 if objective == "energy" else np.ones(count) + energies / (count * 3.2 + 1)
        result = milp(costs, integrality=np.ones(count), bounds=Bounds(0, 1),
                      constraints=LinearConstraint(matrix.tocsr(),
                          [c[1] for c in constraints], [c[2] for c in constraints]),
                      options={"time_limit": 120.0, "mip_rel_gap": 0.001})
        if result.x is not None:
            selected = [dict(c) for c, x in zip(candidates, result.x) if x > 0.5]
    selected.sort(key=lambda c: c["preparation_start_s"])
    drone_ready = {"R01": 0.0, "R02": 0.0}
    battery_ready = {f"RE-{i:02d}": 0.0 for i in range(1, relay.battery_count + 1)}
    for index, item in enumerate(selected, 1):
        start = item["preparation_start_s"]
        drone = next(code for code, ready in drone_ready.items() if ready <= start + 1e-7)
        battery = next(code for code, ready in battery_ready.items() if ready <= start + 1e-7)
        item.update(relay_sortie=f"R3-{index:03d}", relay_drone=drone, energy_component=battery)
        drone_ready[drone] = item["return_o01_s"] + relay.turnaround_s
        battery_ready[battery] = item["return_o01_s"] + _charge_duration(
            relay.full_charge_s, item["return_soc_percent"] / 100)
    covered = {gap for c in selected for gap in c["covered_gap_ids"]}
    uncovered = sorted(set(range(1, gap_count + 1)) - covered)
    return {"feasible_cover": not uncovered, "selected_count": len(selected),
            "selected": selected, "uncovered_gap_ids": uncovered,
            "impossible_gap_ids": impossible, "best_partial_covered_count": len(covered),
            "drone_ready_times": drone_ready, "energy_component_ready_times": battery_ready,
            "objective": objective,
            "solver_status": result.message if result is not None else "empty or impossible candidate set",
            "mip_gap": float(result.mip_gap) if result is not None and result.x is not None else None,
            "method": "binary interval-cover MILP with airframe and component overlap cliques; interval coloring"}


def run(output_dir: Path = OUTPUT_DEFAULT, sample_step_s: float = 10.0,
        relay_candidate_step_s: float = 10.0) -> dict:
    sorties, transport_metrics, boxes = solve_problem2()
    phases = build_trajectory(sorties)
    samples, intervals = screen_direct_links(sorties, phases, sample_step_s)
    gaps = _merge_gap_intervals(intervals)
    _assign_gap_ids(intervals, gaps)
    drones, batteries = load_resources()
    coordinated_transport, transport_coordination = coordinate_transport_starts(
        sorties, boxes, intervals, batteries, step_s=60.0,
    )
    if transport_coordination["delays_s"]:
        sorties = coordinated_transport
        phases = build_trajectory(sorties)
        samples, intervals = screen_direct_links(sorties, phases, sample_step_s)
        gaps = _merge_gap_intervals(intervals)
        _assign_gap_ids(intervals, gaps)
    relay_candidates = search_relay_candidates(sorties, phases, gaps, relay_candidate_step_s)
    from src.problem3.joint import coordinate_joint_schedule
    coordinated_sorties, selected, coordination = coordinate_joint_schedule(sorties, relay_candidates)
    if coordinated_sorties is not None:
        sorties = coordinated_sorties
        phases = build_trajectory(sorties)
        samples, intervals = screen_direct_links(sorties, phases, sample_step_s)
        gaps = _merge_gap_intervals(intervals)
        _assign_gap_ids(intervals, gaps)
    flat_candidates = [{"candidate_group": group["group"], **candidate}
                       for group in relay_candidates for candidate in group["candidates"]]
    selected_schedule = coordination.get("relay_schedule", {})
    coordination["transport_coordination"] = transport_coordination
    coordination["transport_delays_s"] = transport_coordination["delays_s"]
    relay_schedule = {"selected": selected, "selected_count": len(selected),
                      "feasible_cover": coordination["feasible"],
                      "uncovered_gap_ids": list(selected_schedule.get("uncovered_gap_ids", [])),
                      "impossible_gap_ids": list(selected_schedule.get("impossible_gap_ids", [])),
                      "method": "joint transport-delay and relay-interval MILP"}
    transport_evaluator = RouteEvaluator(0.2, 1.0)
    try:
        transport_metrics = _validate(sorties, boxes, drones, batteries, transport_evaluator)
    finally:
        transport_evaluator.close()
    from src.problem3.audit import audit_communication, audit_relay_resources
    communication_rows, communication_metrics = audit_communication(phases, selected)
    relay_metrics = audit_relay_resources(selected)
    hard_deadline_metrics = _hard_deadline_audit(sorties)
    audit_evaluator = LinkEvaluator()
    try:
        elevation_audit = _node_elevation_audit(audit_evaluator)
    finally:
        audit_evaluator.close()
    metrics = {
        "status": "relay candidate/partial schedule diagnostic; not a feasible Q3 solution",
        "transport_feasible": transport_metrics["feasible"],
        "box_count": len(boxes),
        "sortie_count": len(sorties),
        "direct_link_threshold_db": link_limits(load_link_parameters())["transport_gateway_db"],
        "sample_step_s": sample_step_s,
        "transport_start_coordination": coordination,
        "sample_count": len(samples),
        "direct_gap_sample_count": sum(not item["available"] for item in samples),
        "direct_gap_interval_count": len(gaps),
        "relay_candidate_step_s": relay_candidate_step_s,
        "relay_candidates_saved_per_group": "energy, departure and return representatives per coverage subset",
        "relay_candidate_summary": [
            {key: gap[key] for key in ("group", "covered_sorties", "gap_count", "gap_start_s", "gap_end_s", "candidate_count")}
            for gap in relay_candidates
        ],
        "gap_groups_with_relay_candidate": sum(item["candidate_count"] > 0 for item in relay_candidates),
        "relay_candidate_feasibility": "sampled spatial candidates; final feasibility is determined by the independent interval audit",
        "relay_schedule_feasibility": relay_schedule["method"],
        "relay_schedule": {key: value for key, value in relay_schedule.items() if key != "selected"},
        "DSM_node_elevation_audit_count": len(elevation_audit),
        "continuity_certified": communication_metrics["certified"],
        "communication_validation": communication_metrics,
        "transport_validation": transport_metrics,
        "hard_deadline_audit": hard_deadline_metrics,
        "relay_validation": relay_metrics,
        "joint_makespan_s": max([s.return_s for s in sorties] + [r["return_o01_s"] for r in relay_schedule["selected"]]),
        "transport_energy_kwh": sum(s.energy_kwh for s in sorties),
        "relay_energy_kwh": sum(r["energy_kwh"] for r in relay_schedule["selected"]),
        "joint_energy_kwh": sum(s.energy_kwh for s in sorties) + sum(r["energy_kwh"] for r in relay_schedule["selected"]),
        "note": "Candidate screening is sampled. Final communication certification uses conservative swept-cell interval bounds, not sample interpolation.",
    }
    metrics["feasible"] = (transport_metrics["feasible"] and hard_deadline_metrics["passed"]
                           and relay_schedule["feasible_cover"]
                           and communication_metrics["certified"] and relay_metrics["feasible"])
    if metrics["feasible"]:
        metrics["status"] = "feasible joint schedule with interval-certified communication under the stated DSM/LOS model"
    output_dir.mkdir(parents=True, exist_ok=True)
    _write_csv(output_dir / "link_samples.csv", samples)
    _write_csv(output_dir / "direct_link_intervals.csv", intervals)
    _write_csv(output_dir / "relay_candidates.csv", flat_candidates)
    _write_csv(output_dir / "relay_schedule.csv", relay_schedule["selected"])
    _write_csv(output_dir / "node_dsm_elevation_audit.csv", elevation_audit)
    _write_csv(output_dir / "relay_uncovered_gaps.csv", [
        row for row in communication_rows if row["保障方式"] == "中断"])
    _write_csv(output_dir / "gap_id_map.csv", gaps)
    _write_csv(output_dir / "communication_audit.csv",
               communication_rows)
    _write_csv(output_dir / "transport_inherited_audit.csv", [
        {
            "架次编号": sortie.code,
            "无人机编号": sortie.drone,
            "机型编号": sortie.aircraft,
            "电池编号": sortie.battery,
            "准备开始时刻（s）": sortie.prep_start_s,
            "起飞时刻（s）": sortie.takeoff_s,
            "返回O01时刻（s）": sortie.return_s,
            "访问服务区顺序": "→".join(sortie.route),
            "架次能耗（kWh）": sortie.energy_kwh,
            "返航SOC（%）": sortie.battery_soc_return_percent,
            "逐箱交付数": len(sortie.boxes),
        }
        for sortie in sorties
    ])
    _write_csv(output_dir / "relay_resource_audit.csv", [
        {
            "中继架次编号": item.get("relay_sortie", ""),
            "中继无人机编号": item.get("relay_drone", ""),
            "能源组件编号": item.get("energy_component", ""),
            "准备开始时刻（s）": item["preparation_start_s"],
            "建链完成/服务开始（s）": item["service_start_s"],
            "服务结束时刻（s）": item["service_end_s"],
            "返回O01时刻（s）": item["return_o01_s"],
            "机体再次可用（s）": item["return_o01_s"] + load_relay_parameters().turnaround_s,
            "能源组件再次可用（s）": item["return_o01_s"] + _charge_duration(
                load_relay_parameters().full_charge_s, item["return_soc_percent"] / 100),
            "架次能耗（kWh）": item["energy_kwh"],
            "返航SOC（%）": item["return_soc_percent"],
        }
        for item in relay_schedule["selected"]
    ])
    _write_q3_submission(output_dir, sorties, intervals, relay_schedule["selected"], communication_rows)
    _plot_problem3_overview(output_dir / "routes_relays.png", phases, relay_schedule["selected"])
    _package_program(output_dir)
    (output_dir / "screening.json").write_text(json.dumps(metrics, ensure_ascii=False, indent=2), encoding="utf-8")
    return metrics


def main() -> None:
    parser = argparse.ArgumentParser(description="Problem 3 communication and relay scheduler")
    parser.add_argument("--output", type=Path, default=OUTPUT_DEFAULT)
    parser.add_argument("--sample-step", type=float, default=10.0)
    parser.add_argument("--relay-candidate-step", type=float, default=10.0)
    args = parser.parse_args()
    if args.sample_step <= 0 or args.relay_candidate_step <= 0:
        parser.error("sampling steps must be positive")
    print(json.dumps(run(args.output, args.sample_step, args.relay_candidate_step), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
