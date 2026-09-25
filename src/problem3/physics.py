from __future__ import annotations

import math
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import rasterio
from openpyxl import load_workbook

from src.problem1.solver import BASE_DATA, DEM_PATH, Node, great_circle_distance_m, load_nodes
from src.problem21.solver import _rows


@dataclass(frozen=True)
class Radio:
    transmit_dbm: float
    antenna_gain_dbi: float


@dataclass(frozen=True)
class LinkParameters:
    frequency_mhz: float
    system_loss_db: float
    obstruction_loss_db: float
    sensitivity_dbm: float
    fade_margin_db: float
    transport: Radio
    relay_access: Radio
    relay_backhaul: Radio
    gateway: Radio


@dataclass(frozen=True)
class LinkResult:
    distance_3d_m: float
    fspl_db: float | None
    obstruction: bool | None
    path_loss_db: float | None
    threshold_db: float
    available: bool
    reason: str | None = None


@dataclass(frozen=True)
class RelayParameters:
    code: str
    empty_mass_kg: float
    communication_mass_kg: float
    takeoff_mass_kg: float
    cruise_speed_mps: float
    cruise_power_kw: float
    usable_energy_kwh: float
    reserve_fraction: float
    preparation_s: float
    link_setup_s: float
    turnaround_s: float
    climb_speed_mps: float
    descent_speed_mps: float
    climb_efficiency: float
    descent_efficiency: float
    hover_power_kw: float
    communication_power_kw: float
    maximum_agl_m: float
    battery_count: int
    full_charge_s: float


@dataclass(frozen=True)
class RelayMissionEstimate:
    outbound_flight_s: float
    return_flight_s: float
    link_setup_s: float
    service_s: float
    return_time_s: float
    energy_kwh: float
    return_soc_percent: float
    feasible_energy: bool


def load_link_parameters(path: Path | None = None) -> LinkParameters:
    workbook = load_workbook(path or BASE_DATA / "通信链路参数.xlsx", read_only=True, data_only=True)
    rows = list(workbook.active.iter_rows(values_only=True))
    workbook.close()
    values = {(str(row[0]), str(row[3])): float(row[4])
              for row in rows[2:] if row[0] and row[1] and row[3] and isinstance(row[4], (int, float))}

    def parameter(category: str, symbol: str) -> float:
        try:
            return values[(category, symbol)]
        except KeyError as error:
            raise ValueError(f"Missing communication parameter {category}/{symbol}") from error

    def radio(category: str) -> Radio:
        return Radio(parameter(category, "Pt"), parameter(category, "G"))

    return LinkParameters(
        frequency_mhz=parameter("传播参数", "f"),
        system_loss_db=parameter("传播参数", "Lsys"),
        obstruction_loss_db=parameter("传播参数", "Lobs"),
        sensitivity_dbm=parameter("接收参数", "Psens"),
        fade_margin_db=parameter("接收参数", "M"),
        transport=radio("运输无人机"),
        relay_access=radio("中继接入端"),
        relay_backhaul=radio("中继回传端"),
        gateway=radio("固定网关 G01"),
    )


def load_relay_parameters(path: Path | None = None) -> RelayParameters:
    rows = _rows(path or BASE_DATA / "中继无人机数据.xlsx")
    parameters = next((row for row in rows if row[0] == "R" and isinstance(row[4], (int, float))), None)
    inventory = next((row for row in rows if row[0] == "R" and isinstance(row[1], (int, float))), None)
    if parameters is None or inventory is None:
        raise ValueError("Relay aircraft or shared energy-component parameters are missing")
    return RelayParameters(
        code=str(parameters[0]), empty_mass_kg=float(parameters[2]),
        communication_mass_kg=float(parameters[3]), takeoff_mass_kg=float(parameters[4]),
        cruise_speed_mps=float(parameters[5]), cruise_power_kw=float(parameters[6]),
        usable_energy_kwh=float(parameters[7]), reserve_fraction=float(parameters[8]) / 100,
        preparation_s=float(parameters[9]), link_setup_s=float(parameters[10]),
        turnaround_s=float(parameters[11]), climb_speed_mps=float(parameters[12]),
        descent_speed_mps=float(parameters[13]), climb_efficiency=float(parameters[14]),
        descent_efficiency=float(parameters[15]), hover_power_kw=float(parameters[16]),
        communication_power_kw=float(parameters[17]), maximum_agl_m=float(parameters[18]),
        battery_count=int(inventory[1]), full_charge_s=float(inventory[2]),
    )


def estimate_relay_mission(service_s: float, outbound_distance_m: float, outbound_max_dsm_m: float,
                           hover_altitude_m: float, return_distance_m: float, return_max_dsm_m: float,
                           base_altitude_m: float, parameters: RelayParameters | None = None) -> RelayMissionEstimate:
    params = parameters or load_relay_parameters()
    if service_s < 0 or min(outbound_distance_m, return_distance_m) < 0:
        raise ValueError("Relay mission durations and distances must be nonnegative")

    def leg(distance_m: float, terrain_max_m: float, start_m: float, end_m: float) -> tuple[float, float]:
        cruise_altitude_m = max(terrain_max_m + 50, start_m, end_m)
        climb_m = max(0.0, cruise_altitude_m - start_m)
        descent_m = max(0.0, cruise_altitude_m - end_m)
        climb_s = climb_m / params.climb_speed_mps
        cruise_s = distance_m / params.cruise_speed_mps
        descent_s = descent_m / params.descent_speed_mps
        horizontal_kwh = params.cruise_power_kw * cruise_s / 3600
        climb_kwh = params.takeoff_mass_kg * 9.80665 * climb_m / (3_600_000 * params.climb_efficiency)
        descent_kwh = 0.0
        if params.descent_efficiency > 0:
            descent_kwh = params.takeoff_mass_kg * 9.80665 * descent_m / (3_600_000 * params.descent_efficiency)
        return climb_s + cruise_s + descent_s, horizontal_kwh + climb_kwh + descent_kwh

    outbound_s, outbound_kwh = leg(outbound_distance_m, outbound_max_dsm_m, base_altitude_m, hover_altitude_m)
    return_s, return_kwh = leg(return_distance_m, return_max_dsm_m, hover_altitude_m, base_altitude_m)
    hover_power = params.hover_power_kw + params.communication_power_kw
    hover_kwh = hover_power * (params.link_setup_s + service_s) / 3600
    energy = outbound_kwh + return_kwh + hover_kwh
    return RelayMissionEstimate(
        outbound_flight_s=outbound_s, return_flight_s=return_s,
        link_setup_s=params.link_setup_s, service_s=service_s,
        return_time_s=params.preparation_s + outbound_s + params.link_setup_s + service_s + return_s,
        energy_kwh=energy,
        return_soc_percent=100 * (1 - energy / params.usable_energy_kwh),
        feasible_energy=energy <= params.usable_energy_kwh * (1 - params.reserve_fraction) + 1e-10,
    )


def directional_loss_limit_db(tx: Radio, rx: Radio, params: LinkParameters) -> float:
    receive_threshold = params.sensitivity_dbm + params.fade_margin_db
    return tx.transmit_dbm + tx.antenna_gain_dbi + rx.antenna_gain_dbi - params.system_loss_db - receive_threshold


def bidirectional_limit_db(first: Radio, second: Radio, params: LinkParameters) -> float:
    return min(directional_loss_limit_db(first, second, params),
               directional_loss_limit_db(second, first, params))


def link_limits(params: LinkParameters) -> dict[str, float]:
    return {
        "transport_gateway_db": bidirectional_limit_db(params.transport, params.gateway, params),
        "transport_relay_db": bidirectional_limit_db(params.transport, params.relay_access, params),
        "relay_gateway_db": bidirectional_limit_db(params.relay_backhaul, params.gateway, params),
    }


def _segment_elevations(dem: rasterio.io.DatasetReader, first: Node, second: Node,
                        first_altitude_m: float, second_altitude_m: float,
                        step_m: float = 30.0, values=None) -> tuple[float, ...] | None:
    distance = great_circle_distance_m(first, second)
    intervals = max(1, math.ceil(distance / step_m))
    fractions = np.arange(1, intervals, dtype=float) / intervals
    longitude = first.longitude + fractions * (second.longitude - first.longitude)
    latitude = first.latitude + fractions * (second.latitude - first.latitude)
    rows, cols = rasterio.transform.rowcol(dem.transform, longitude, latitude)
    rows, cols = np.asarray(rows), np.asarray(cols)
    if np.any((rows < 0) | (rows >= dem.height) | (cols < 0) | (cols >= dem.width)):
        return None
    raster = values if values is not None else dem.read(1)
    heights = raster[rows, cols]
    invalid = ~np.isfinite(heights) | np.isclose(heights, -32767.0)
    if dem.nodata is not None:
        invalid |= np.isclose(heights, dem.nodata)
    if np.any(invalid):
        return None
    return heights - (first_altitude_m + fractions * (second_altitude_m - first_altitude_m))



class LinkEvaluator:
    def __init__(self, params: LinkParameters | None = None, dem_path: Path = DEM_PATH):
        self.params = params or load_link_parameters()
        self.dem = rasterio.open(dem_path)
        self.dem_values = self.dem.read(1)
        self.base, self.sites = load_nodes()

    def close(self) -> None:
        self.dem.close()

    def evaluate(self, first: Node, first_altitude_m: float, second: Node, second_altitude_m: float,
                 threshold_db: float) -> LinkResult:
        horizontal_m = great_circle_distance_m(first, second)
        distance_m = math.hypot(horizontal_m, first_altitude_m - second_altitude_m)
        if distance_m <= 0 and (first.code == "G01" or second.code == "G01"):
            return LinkResult(0.0, None, False, None, threshold_db, True,
                              "co-located gateway: zero-distance continuous extension")
        if distance_m <= 0:
            return LinkResult(0.0, None, None, None, threshold_db, False, "coincident 3D endpoints")
        terrain_clearances = _segment_elevations(
            self.dem, first, second, first_altitude_m, second_altitude_m,
            values=self.dem_values,
        )
        if terrain_clearances is None:
            return LinkResult(distance_m, None, None, None, threshold_db, False, "DEM NoData or outside extent")
        obstructed = any(clearance >= 0 for clearance in terrain_clearances)
        frequency_term = 20 * math.log10(self.params.frequency_mhz)
        fspl_db = 32.45 + frequency_term + 20 * math.log10(distance_m / 1000)
        path_loss_db = fspl_db + (self.params.obstruction_loss_db if obstructed else 0.0)
        return LinkResult(distance_m, fspl_db, obstructed, path_loss_db, threshold_db,
                          path_loss_db <= threshold_db)


def certify_moving_link(evaluator: LinkEvaluator, first: Node, first_altitude: float,
                        last: Node, last_altitude: float, fixed: Node, fixed_altitude: float,
                        threshold_db: float) -> bool:
    """Conservative interval bound for linear motion and the 30 m LOS rule.

    Enumerate all possible LOS sample counts. Bound each moving sample by
    every raster cell in its swept coordinate rectangle and minimum LOS height.
    False means unproved, not necessarily an outage. Subdivide and retry.
    """
    if fixed.code != "G01":
        origin = np.array([first.longitude - fixed.longitude, first.latitude - fixed.latitude,
                           first_altitude - fixed_altitude])
        delta = np.array([last.longitude - first.longitude, last.latitude - first.latitude,
                          last_altitude - first_altitude])
        norm = float(delta @ delta)
        fraction = min(1.0, max(0.0, -float(origin @ delta) / norm)) if norm else 0.0
        if np.linalg.norm(origin + fraction * delta) < 1e-10:
            return False
    radius = 6371008.8
    movement = radius * math.hypot(math.radians(last.latitude - first.latitude),
                                 math.radians(last.longitude - first.longitude))
    middle = Node("mid", (first.longitude + last.longitude) / 2,
                  (first.latitude + last.latitude) / 2, 0)
    horizontal = great_circle_distance_m(middle, fixed)
    max_distance = math.hypot(horizontal + movement / 2,
                             max(abs(first_altitude - fixed_altitude),
                                 abs(last_altitude - fixed_altitude)))
    if max_distance <= 0:
        return False
    min_count = max(1, math.ceil(max(0, horizontal - movement / 2) / 30))
    max_count = max(1, math.ceil((horizontal + movement / 2) / 30))
    if max_count - min_count > 4:
        return False
    terrain_clear = True
    inverse = ~evaluator.dem.transform
    for count in range(min_count, max_count + 1):
        fractions = np.arange(1, count, dtype=float) / count
        if not len(fractions):
            continue
        x0 = first.longitude + fractions * (fixed.longitude - first.longitude)
        y0 = first.latitude + fractions * (fixed.latitude - first.latitude)
        x1 = last.longitude + fractions * (fixed.longitude - last.longitude)
        y1 = last.latitude + fractions * (fixed.latitude - last.latitude)
        col0, row0 = inverse * (x0, y0)
        col1, row1 = inverse * (x1, y1)
        low_row = np.floor(np.minimum(row0, row1) - 1e-9).astype(int)
        high_row = np.floor(np.maximum(row0, row1) + 1e-9).astype(int)
        low_col = np.floor(np.minimum(col0, col1) - 1e-9).astype(int)
        high_col = np.floor(np.maximum(col0, col1) + 1e-9).astype(int)
        if np.any((low_row < 0) | (high_row >= evaluator.dem.height)
                  | (low_col < 0) | (high_col >= evaluator.dem.width)):
            return False
        max_rows = int(np.max(high_row - low_row))
        max_cols = int(np.max(high_col - low_col))
        if max_rows > 8 or max_cols > 8:
            return False
        heights = np.full(len(fractions), -np.inf)
        for dr in range(max_rows + 1):
            for dc in range(max_cols + 1):
                rows = np.minimum(low_row + dr, high_row)
                cols = np.minimum(low_col + dc, high_col)
                values = evaluator.dem_values[rows, cols]
                invalid = ~np.isfinite(values) | np.isclose(values, -32767)
                if evaluator.dem.nodata is not None:
                    invalid |= np.isclose(values, evaluator.dem.nodata)
                if np.any(invalid):
                    return False
                heights = np.maximum(heights, values)
        minimum_los = min(first_altitude, last_altitude) * (1 - fractions) + fixed_altitude * fractions
        if np.any(heights >= minimum_los - 1e-8):
            terrain_clear = False
    upper_loss = (32.45 + 20 * math.log10(evaluator.params.frequency_mhz)
                  + 20 * math.log10(max_distance / 1000)
                  + (0 if terrain_clear else evaluator.params.obstruction_loss_db))
    return upper_loss <= threshold_db - 1e-9


def sampled_flight_leg(evaluator: LinkEvaluator, first: Node, last: Node) -> tuple[float, float]:
    """Same 30 m endpoint-inclusive samples as Q2, using the loaded raster."""
    distance = great_circle_distance_m(first, last)
    fractions = np.linspace(0, 1, max(1, math.ceil(distance / 30)) + 1)
    rows, cols = rasterio.transform.rowcol(
        evaluator.dem.transform,
        first.longitude + fractions * (last.longitude - first.longitude),
        first.latitude + fractions * (last.latitude - first.latitude))
    rows, cols = np.asarray(rows), np.asarray(cols)
    if np.any((rows < 0) | (rows >= evaluator.dem.height) | (cols < 0) | (cols >= evaluator.dem.width)):
        raise ValueError("Relay flight outside DEM")
    values = evaluator.dem_values[rows, cols]
    invalid = ~np.isfinite(values) | np.isclose(values, -32767)
    if evaluator.dem.nodata is not None:
        invalid |= np.isclose(values, evaluator.dem.nodata)
    if np.any(invalid):
        raise ValueError("Relay flight crosses DEM NoData")
    return float(np.max(values)), distance
