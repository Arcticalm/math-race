from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import shutil
from collections import defaultdict
from functools import lru_cache
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT / "data" / "无人机应急物资运输基础数据"
DEFAULT_INPUT = ROOT / "outputs" / "q3"
DEFAULT_OUTPUT = ROOT / "outputs" / "q4"
SITES = tuple(f"S{i:03d}" for i in range(1, 16))
AIR_TYPES = ("A", "B", "C")
RESOURCE_KEYS = ("A_airframe", "B_airframe", "C_airframe", "A_battery", "B_battery", "C_battery", "relay_airframe", "relay_energy")


def _csv(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def _float(row: dict, *names: str, default: float = 0.0) -> float:
    for name in names:
        if name in row and row[name] not in (None, ""):
            return float(row[name])
    return default


def _number(row: dict, name: str) -> float:
    value = row.get(name)
    if value in (None, ""):
        raise ValueError(f"必需数值字段为空: {name}")
    number = float(value)
    if not math.isfinite(number):
        raise ValueError(f"数值字段非有限: {name}={value}")
    return number


def _sites(value: str) -> tuple[str, ...]:
    return tuple(sorted(set(x.strip() for x in value.replace("→", ",").split(",") if x.strip())))


def load_inventory() -> dict[str, int]:
    """Read counts from the source workbooks; constants are never used as fallback."""
    path = DATA / "运输无人机数据.xlsx"
    rows = list(load_workbook(path, read_only=True, data_only=True).active.iter_rows(values_only=True))
    counts = {f"{kind}_airframe": 0 for kind in AIR_TYPES}
    batteries = {f"{kind}_battery": 0 for kind in AIR_TYPES}
    section = None
    for row in rows:
        if row[0] == "逐架无人机清单":
            section = "drones"
        elif row[0] in AIR_TYPES and isinstance(row[1], (int, float)):
            batteries[f"{row[0]}_battery"] = int(row[1])
        elif section == "drones" and row[0] and str(row[0]).startswith("U"):
            counts[f"{row[1]}_airframe"] += 1
    result = {**counts, **batteries}
    relay_rows = list(load_workbook(DATA / "中继无人机数据.xlsx", read_only=True, data_only=True).active.iter_rows(values_only=True))
    relay_count = sum(1 for row in relay_rows if row[0] and str(row[0]).startswith("R0"))
    result["relay_airframe"] = relay_count
    # Energy component inventory is the explicit component list in the source workbook.
    for row in relay_rows:
        if row[0] == "R" and isinstance(row[1], (int, float)):
            result["relay_energy"] = int(row[1])
            break
    if not result.get("relay_energy"):
        raise ValueError("中继能源组件库存缺失，不能据猜测填补库存")
    return result


def q3_gate(input_dir: Path) -> dict:
    screening_path = input_dir / "screening.json"
    if not screening_path.exists():
        return {"ready": False, "reason": "缺少 Q3 screening.json"}
    screening = json.loads(screening_path.read_text(encoding="utf-8"))
    required = (input_dir / "transport_inherited_audit.csv", input_dir / "relay_schedule.csv", input_dir / "relay_resource_audit.csv", input_dir / "communication_audit.csv")
    missing = [str(path.name) for path in required if not path.exists()]
    if missing:
        return {"ready": False, "reason": f"缺少 Q3 输出: {', '.join(missing)}", "screening": screening}
    certificate = screening.get("continuous_communication_validation", {})
    if certificate.get("audit_version") != "exact_dem_interval_v2":
        return {"ready": False, "reason": "Q3 缺少 exact_dem_interval_v2 连续通信审计证书；旧采样认证不可继承", "screening": screening}
    if screening.get("feasible") is not True or screening.get("continuity_certified") is not True:
        return {"ready": False, "reason": "Q3 联合方案或通信连续性审计未通过", "screening": screening}
    for name in ("transport_inherited_audit.csv", "communication_audit.csv"):
        rows = _csv(input_dir / name)
        if name != "relay_schedule.csv" and not rows:
            return {"ready": False, "reason": f"Q3 审计文件为空: {name}", "screening": screening}
    transport = _csv(input_dir / "transport_inherited_audit.csv")
    communication = _csv(input_dir / "communication_audit.csv")
    if not {"运输架次编号", "保障方式"}.issubset(communication[0]):
        return {"ready": False, "reason": "通信审计缺少架次或保障状态字段", "screening": screening}
    transport_ids = {row.get("架次编号", "") for row in transport}
    communication_ids = {row.get("运输架次编号", "") for row in communication}
    if not transport_ids.issubset(communication_ids):
        return {"ready": False, "reason": "通信审计未覆盖全部运输架次", "screening": screening}
    if any(row.get("保障方式") == "中断" for row in communication):
        return {"ready": False, "reason": "通信审计含中断区间，与 Q3 可行标记矛盾", "screening": screening}
    declared_relays = {row.get("relay_sortie", ""): set(_sites(
        row.get("covered_sorties") or row.get("保障运输架次") or ""))
        for row in _csv(input_dir / "relay_schedule.csv")}
    for row in communication:
        mode = row.get("保障方式")
        if mode not in ("直连", "中继"):
            return {"ready": False, "reason": "Q3 通信审计包含未知保障状态", "screening": screening}
        if mode == "中继" and row.get("运输架次编号") not in declared_relays.get(row.get("中继架次编号"), set()):
            return {"ready": False, "reason": "Q3 实际通信保障关系与中继任务映射不一致", "screening": screening}
    return {"ready": True, "screening": screening}


def validate_frozen_partition(input_dir: Path, result: dict) -> None:
    """Reject stale policy or changed source records before comparing/publishing."""
    if result.get("partition_policy") != PARTITION_POLICY:
        raise ValueError("Q4 must be recomputed with the frozen relay closure policy")
    hashes = result.get("frozen_input_sha256", {})
    required = {"screening.json", "transport_inherited_audit.csv", "relay_schedule.csv",
                "relay_resource_audit.csv", "communication_audit.csv"}
    if not required.issubset(hashes):
        raise ValueError("Incomplete frozen Q3 fingerprints")
    for name, expected in hashes.items():
        path = input_dir / name
        if Path(name).name != name or not path.is_file() or hashlib.sha256(path.read_bytes()).hexdigest() != expected:
            raise ValueError(f"Frozen Q3 input changed: {name}")


def read_q3(input_dir: Path) -> tuple[list[dict], list[dict]]:
    transport = _csv(input_dir / "transport_inherited_audit.csv")
    relay = _csv(input_dir / "relay_schedule.csv")
    resource_rows = _csv(input_dir / "relay_resource_audit.csv")
    relay_resources = {row.get("中继架次编号", ""): row for row in resource_rows}
    relay_ids = [row.get("relay_sortie", "") for row in relay]
    if len(relay_resources) != len(resource_rows) or len(set(relay_ids)) != len(relay_ids):
        raise ValueError("Q3 中继架次编号重复")
    required_transport = {"架次编号", "机型编号", "准备开始时刻（s）", "返回O01时刻（s）", "访问服务区顺序", "返航SOC（%）"}
    if not transport or not required_transport.issubset(transport[0]) or any(
        not _sites(row.get("访问服务区顺序", ""))
        or not set(_sites(row.get("访问服务区顺序", ""))).issubset(SITES)
        or row.get("机型编号") not in AIR_TYPES for row in transport
    ):
        raise ValueError("Q3 运输审计缺少服务区访问关系")
    for row in transport:
        start = _number(row, "准备开始时刻（s）")
        takeoff = _number(row, "起飞时刻（s）")
        end = _number(row, "返回O01时刻（s）")
        soc = _number(row, "返航SOC（%）")
        _number(row, "架次能耗（kWh）")
        if not start <= takeoff < end or not 0 <= soc <= 100:
            raise ValueError(f"Q3 运输时间区间无效: {row.get('架次编号')}")
    if relay:
        for row in relay:
            if "relay_sortie" not in row:
                raise ValueError("Q3 中继排程缺少 relay_sortie 编号")
            resource = relay_resources.get(row.get("relay_sortie", ""))
            if resource is None:
                raise ValueError(f"中继资源审计缺少任务: {row.get('relay_sortie')}")
            row.update(resource)
            required_relay = {"中继架次编号", "准备开始时刻（s）", "建链完成/服务开始（s）", "服务结束时刻（s）", "返回O01时刻（s）", "机体再次可用（s）", "能源组件再次可用（s）", "返航SOC（%）"}
            if not required_relay.issubset(row):
                raise ValueError(f"Q3 中继资源审计缺少字段: {sorted(required_relay - set(row))}")
            prep, service_start = _number(row, "准备开始时刻（s）"), _number(row, "建链完成/服务开始（s）")
            service_end, returned = _number(row, "服务结束时刻（s）"), _number(row, "返回O01时刻（s）")
            ready_airframe, ready_energy = _number(row, "机体再次可用（s）"), _number(row, "能源组件再次可用（s）")
            soc = _number(row, "返航SOC（%）")
            _number(row, "架次能耗（kWh）")
            if not prep <= service_start <= service_end <= returned <= ready_airframe or ready_energy < returned or not 0 <= soc <= 100:
                raise ValueError(f"Q3 中继资源时间或 SOC 无效: {row.get('中继架次编号')}")
    if set(row.get("relay_sortie", "") for row in relay) != set(relay_resources):
        raise ValueError("中继排程与中继资源审计的任务集合不一致")
    assigned_sites = {site for row in transport for site in _sites(row["访问服务区顺序"])}
    if assigned_sites != set(SITES):
        raise ValueError(f"运输任务服务区覆盖不完整: missing={sorted(set(SITES)-assigned_sites)}, unknown={sorted(assigned_sites-set(SITES))}")
    sortie_ids = [row["架次编号"] for row in transport]
    if len(sortie_ids) != len(set(sortie_ids)):
        raise ValueError("Q3 运输架次编号重复")
    known_sorties = set(sortie_ids)
    for row in relay:
        value = row.get("covered_sorties") or row.get("保障运输架次") or ""
        covered = {item.strip() for item in value.replace("→", ",").split(",") if item.strip()}
        if not covered or covered - known_sorties:
            raise ValueError(f"中继任务保障架次映射缺失或未知: {row.get('relay_sortie')}, {sorted(covered-known_sorties)}")
    deliveries = _csv(input_dir / "box_delivery_audit.csv")
    if deliveries:
        by_sortie = {row["架次编号"]: row for row in transport}
        seen_boxes = set()
        for row in transport:
            row["交付质量（kg）"] = 0.0
        counts = defaultdict(int)
        for delivery in deliveries:
            box = delivery.get("货箱编号", "")
            task = delivery.get("架次编号", "")
            if not box or box in seen_boxes or task not in by_sortie:
                raise ValueError("Q3 逐箱交付映射存在重复货箱或未知架次")
            if delivery.get("服务区编号") not in _sites(by_sortie[task]["访问服务区顺序"]):
                raise ValueError(f"Q3 货箱服务区与运输路线不一致: {box}")
            mass = _number(delivery, "质量（kg）")
            if mass <= 0:
                raise ValueError(f"Q3 货箱质量无效: {box}")
            by_sortie[task]["交付质量（kg）"] += mass
            counts[task] += 1
            seen_boxes.add(box)
        if any(counts[row["架次编号"]] != _number(row, "逐箱交付数") for row in transport):
            raise ValueError("Q3 逐箱交付明细数量与运输审计不一致")
    return transport, relay


PARTITION_POLICY = "frozen_transport_and_relay_closure_v1"


def atomic_units(transport: list[dict], relay=()) -> list[frozenset[str]]:
    """Merge the frozen transport/relay task graph into indivisible components.

    A relay mission cannot be duplicated, split or assigned across task groups.
    All sites of its covered transport sorties must therefore share one owner.
    """
    parent = {site: site for site in SITES}
    def find(x):
        while parent[x] != x:
            parent[x] = parent[parent[x]]
            x = parent[x]
        return x
    def union(a, b):
        a, b = find(a), find(b)
        if a != b: parent[b] = a
    for row in transport:
        sites = _sites(row["访问服务区顺序"])
        for site in sites[1:]: union(sites[0], site)
    by_sortie = {row["架次编号"]: _sites(row["访问服务区顺序"]) for row in transport}
    for row in relay:
        covered = _covered_sorties(row)
        if not covered or set(covered) - by_sortie.keys():
            raise ValueError(f"中继架次运输映射缺失或未知: {_relay_id(row)}")
        sites = [site for task in covered for site in by_sortie[task]]
        for site in sites[1:]: union(sites[0], site)
    groups = defaultdict(set)
    for site in SITES: groups[find(site)].add(site)
    return sorted((frozenset(group) for group in groups.values()), key=lambda group: min(group))


def _relay_id(row: dict) -> str:
    return row.get("中继架次编号") or row.get("relay_sortie") or "relay"


def _covered_sorties(row: dict) -> list[str]:
    value = row.get("covered_sorties") or row.get("保障运输架次") or ""
    return [item.strip() for item in value.replace("→", ",").split(",") if item.strip()]


def group_attribution(group, transport: list[dict], relay: list[dict]) -> tuple[set[str], int]:
    """Assign frozen missions to one group; reject any cross-group relay."""
    members = set(group)
    known = {row["架次编号"] for row in transport}
    own: set[str] = set()
    for row in transport:
        sites = set(_sites(row["访问服务区顺序"]))
        if sites <= members:
            own.add(row["架次编号"])
        elif sites & members:
            raise ValueError(f"运输架次跨组: {row['架次编号']}")
    shared = 0
    for row in relay:
        task = _relay_id(row)
        covered = _covered_sorties(row)
        if not covered:
            raise ValueError(f"中继架次缺少运输映射: {task}")
        unknown = set(covered) - known
        if unknown:
            raise ValueError(f"中继架次引用未知运输架次: {sorted(unknown)}")
        inside = [item for item in covered if item in own]
        if not inside:
            continue
        if len(inside) != len(covered):
            raise ValueError(f"中继架次跨组，不能复制或拆分冻结任务: {task}")
        own.add(task)
    return own, shared


def partitions(units: list[frozenset[str]], k: int, max_candidates: int | None = None):
    """Enumerate each unlabelled nonempty partition once; no default truncation."""
    if k < 1:
        raise ValueError("任务组数量必须为正整数")
    if len(units) < k:
        return
    count = 0

    def visit(index, groups):
        nonlocal count
        if len(groups) + len(units) - index < k:
            return
        if index == len(units):
            if len(groups) == k:
                count += 1
                if max_candidates is not None and count > max_candidates:
                    raise RuntimeError(f"分区候选超过规模上限 {max_candidates}，无法进行有保证的穷举")
                yield tuple(frozenset().union(*group) for group in groups)
            return
        unit = units[index]
        for group in groups:
            group.append(unit)
            yield from visit(index + 1, groups)
            group.pop()
        if len(groups) < k:
            groups.append([unit])
            yield from visit(index + 1, groups)
            groups.pop()

    yield from visit(0, [])


def _peak(intervals: list[tuple[float, float]]) -> int:
    events = [(start, 1) for start, end in intervals if end > start] + [(end, -1) for start, end in intervals if end > start]
    active = peak = 0
    for _, change in sorted(events, key=lambda item: (item[0], item[1])):
        active += change; peak = max(peak, active)
    return peak


def _charge_s(soc_percent: float, full_s: float) -> float:
    s = max(0.0, min(1.0, soc_percent / 100.0))
    return full_s * (0.65 * (0.90 - s) / 0.90 + 0.35 if s < 0.90 else 0.35 * (1 - s) / 0.10)


@lru_cache(maxsize=1)
def _battery_full_charge_s() -> dict[str, float]:
    full_charge_s = {}
    workbook = load_workbook(DATA / "运输无人机数据.xlsx", read_only=True, data_only=True)
    for values_row in workbook.active.iter_rows(values_only=True):
        if values_row[0] in AIR_TYPES and isinstance(values_row[2], (int, float)):
            full_charge_s[str(values_row[0])] = float(values_row[2])
    workbook.close()
    if set(full_charge_s) != set(AIR_TYPES):
        raise ValueError("运输无人机工作簿缺少机型充电时长参数")
    return full_charge_s


def resource_intervals(transport: list[dict], relay: list[dict]) -> dict[str, list[tuple[float, float, str]]]:
    values = defaultdict(list)
    full_charge_s = _battery_full_charge_s()
    for row in transport:
        kind = row["机型编号"]
        start, end = _float(row, "准备开始时刻（s）", "开始时刻（s）"), _float(row, "返回O01时刻（s）")
        values[f"{kind}_airframe"].append((start, end, row["架次编号"]))
        soc = _float(row, "返航SOC（%）", default=0.0)
        full = full_charge_s[kind]
        charged = end + _charge_s(soc, full)
        if row.get("电池充电完成时刻（s）") not in (None, "") and abs(
            _number(row, "电池充电完成时刻（s）") - charged
        ) > 1e-6:
            raise ValueError(f"Q3 电池充电完成时刻不符合两阶段充电规则: {row['架次编号']}")
        values[f"{kind}_battery"].append((start, charged, row["架次编号"]))
    for row in relay:
        start = _float(row, "准备开始时刻（s）")
        end = _float(row, "机体再次可用（s）")
        values["relay_airframe"].append((start, end, row.get("中继架次编号", row.get("relay_sortie", "relay"))))
        component_end = _float(row, "能源组件再次可用（s）")
        values["relay_energy"].append((start, component_end, row.get("中继架次编号", "relay")))
    return values


def evaluate_partition(groups, transport, relay, inventory, intervals=None, group_cache=None):
    if not groups or any(not group for group in groups) or set().union(*groups) != set(SITES) or sum(map(len, groups)) != len(SITES):
        raise ValueError("每个服务区必须且只能属于一个非空任务组")
    if group_cache is not None and all(group in group_cache for group in groups):
        return [{**group_cache[group], "group": index} for index, group in enumerate(groups, 1)]
    intervals = intervals if intervals is not None else resource_intervals(transport, relay)
    rows = []
    tasks_by_group = {index: group_attribution(group, transport, relay)
                      for index, group in enumerate(groups, 1)}
    for index, group in enumerate(groups, 1):
        if group_cache is not None and group in group_cache:
            rows.append({**group_cache[group], "group": index})
            continue
        own, shared = tasks_by_group[index]
        demand = {}
        for key in RESOURCE_KEYS:
            demand[key] = _peak([(a, b) for a, b, task in intervals[key] if task in own])
        work_sorties = [row for row in transport if row["架次编号"] in own]
        relay_sorties = [row for row in relay if _relay_id(row) in own]
        rows.append({"group": index, "sites": ",".join(sorted(group)), **demand,
                     "transport_sorties": ",".join(sorted(row["架次编号"] for row in work_sorties)),
                     "relay_sorties": ",".join(sorted(_relay_id(row) for row in relay_sorties)),
                     "shared_relay_missions": shared,
                     "site_count": len(group), "sortie_count": len(work_sorties),
                     "transport_flight_time_s": sum(_float(row, "返回O01时刻（s）") - _float(row, "起飞时刻（s）") for row in work_sorties),
                     "box_count": sum(int(_float(row, "逐箱交付数")) for row in work_sorties),
                     "relay_time_s": sum(_float(row, "服务结束时刻（s）") - _float(row, "建链完成/服务开始（s）") for row in relay_sorties),
                     "transport_energy_kwh": sum(_float(row, "架次能耗（kWh）") for row in work_sorties),
                     "relay_energy_kwh": sum(_float(row, "架次能耗（kWh）") for row in relay_sorties)})
        if all("交付质量（kg）" in row for row in work_sorties):
            rows[-1]["cargo_mass_kg"] = sum(_float(row, "交付质量（kg）") for row in work_sorties)
        if group_cache is not None:
            group_cache[group] = {key: value for key, value in rows[-1].items() if key != "group"}
    return rows


def _cv(values: list[float]) -> float | None:
    mean = sum(values) / len(values) if values else 0.0
    return None if mean == 0 else math.sqrt(sum((x - mean) ** 2 for x in values) / len(values)) / mean


def _partition_score(rows: list[dict], inventory: dict[str, int]) -> tuple:
    demand = {key: sum(row[key] for row in rows) for key in RESOURCE_KEYS}
    missing = sum(max(0, demand[key] - inventory[key]) for key in RESOURCE_KEYS)
    total = sum(demand.values())
    cv = max((_cv([row[field] for row in rows]) or 0.0 for field in
              ("transport_flight_time_s", "relay_time_s", "sortie_count", "box_count")), default=0.0)
    return missing, total, cv


def _write_csv(path: Path, rows: list[dict], fields: list[str] | None = None) -> None:
    if fields is None:
        fields = list(dict.fromkeys(key for row in rows for key in row))
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def _write_submission_workbook(path: Path, rows: list[dict]) -> None:
    """Fill the Q4 sheet of the official submission template."""
    source = ROOT / "docs" / "结果提交模板.xlsx"
    shutil.copy2(source, path)
    workbook = load_workbook(path)
    sheet = workbook["Q4_分区配置"]
    headers = [sheet.cell(1, column).value for column in range(1, 12)]
    for row_index in range(2, sheet.max_row + 1):
        for column in range(1, 12):
            sheet.cell(row_index, column).value = None
    for row_index, row in enumerate(rows, 2):
        for column, header in enumerate(headers, 1):
            key = "K" if header == "K（2或3）" else header
            sheet.cell(row_index, column).value = row.get(key, "")
    workbook.save(path)


def run(input_dir: Path = DEFAULT_INPUT, output_dir: Path = DEFAULT_OUTPUT) -> dict:
    output_dir.mkdir(parents=True, exist_ok=True)
    gate = q3_gate(input_dir)
    if not gate["ready"]:
        result = {"status": "等待 Q3 最终方案", "q3_gate": gate}
        (output_dir / "status.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
        return result
    transport, relay = read_q3(input_dir)
    inventory = load_inventory()
    units = atomic_units(transport, relay)
    frozen_files = ("screening.json", "transport_inherited_audit.csv", "relay_schedule.csv",
                    "relay_resource_audit.csv", "communication_audit.csv", "box_delivery_audit.csv")
    result = {"status": "ready", "partition_policy": PARTITION_POLICY, "inventory": inventory, "atomic_units": [sorted(x) for x in units],
              "frozen_input_sha256": {name: hashlib.sha256((input_dir / name).read_bytes()).hexdigest()
                                      for name in frozen_files if (input_dir / name).exists()},
              "solutions": {}}
    interval_data = resource_intervals(transport, relay)
    for key in RESOURCE_KEYS:
        if _peak([(a, b) for a, b, _ in interval_data[key]]) > inventory[key]:
            raise ValueError(f"Q3 冻结排程的全局占用已超过库存: {key}")
    group_cache = {}
    template_rows, inventory_rows, workload_rows, mapping_rows, interval_rows, site_mapping_rows = [], [], [], [], [], []
    for k in (2, 3):
        best = None
        candidate_count = 0
        for candidate in partitions(units, k):
            candidate_count += 1
            rows = evaluate_partition(candidate, transport, relay, inventory, interval_data, group_cache)
            score = _partition_score(rows, inventory)
            if best is None or score < best[0]:
                best = (score, candidate, rows)
        if best is None:
            result["solutions"][str(k)] = {
                "feasible": False, "candidate_count": candidate_count,
                "reason": "原子任务单元数量小于 K"}
            continue
        score, groups, rows = best
        relay_missions = {_relay_id(row) for row in relay}
        duplicated = sorted(task for task in relay_missions if sum(
            task in row["relay_sorties"].split(",") for row in rows) > 1)
        result["solutions"][str(k)] = {"feasible": True, "candidate_count": candidate_count,
                                          "inventory_feasible": score[0] == 0,
                                          "partition_search_complete": True,
                                          "shared_relay_missions": duplicated,
                                          "optimality_scope": "all partitions of the frozen Q3 transport-and-relay task graph under the stated lexicographic objective; every mission has exactly one owner",
                                          "selection_rule": "minimize total inventory deficit, then total resource demand, then max workload CV",
                                          "score": score, "groups": rows}
        for row in rows:
            template_rows.append({"K": k, "任务组编号": f"K{k}-G{row['group']:02d}", "服务区列表": row["sites"],
                                  "A型运输无人机数": row["A_airframe"], "B型运输无人机数": row["B_airframe"], "C型运输无人机数": row["C_airframe"],
                                  "A型电池组数": row["A_battery"], "B型电池组数": row["B_battery"], "C型电池组数": row["C_battery"],
                                  "中继无人机数": row["relay_airframe"], "中继能源组件数": row["relay_energy"]})
            workload_rows.append({"K": k, **row, "transport_flight_time_mean_s": sum(x["transport_flight_time_s"] for x in rows) / k,
                                  "transport_flight_time_min_s": min(x["transport_flight_time_s"] for x in rows),
                                  "transport_flight_time_max_s": max(x["transport_flight_time_s"] for x in rows),
                                  "transport_flight_time_cv": _cv([x["transport_flight_time_s"] for x in rows]),
                                  "relay_service_time_mean_s": sum(x["relay_time_s"] for x in rows) / k,
                                  "relay_service_time_min_s": min(x["relay_time_s"] for x in rows),
                                  "relay_service_time_max_s": max(x["relay_time_s"] for x in rows),
                                  "relay_service_time_cv": _cv([x["relay_time_s"] for x in rows]),
                                  "sortie_count_mean": sum(x["sortie_count"] for x in rows) / k,
                                  "sortie_count_min": min(x["sortie_count"] for x in rows),
                                  "sortie_count_max": max(x["sortie_count"] for x in rows),
                                  "sortie_count_cv": _cv([x["sortie_count"] for x in rows])})
            for field in ("box_count", "transport_energy_kwh", "relay_energy_kwh", "cargo_mass_kg"):
                if all(field in item for item in rows):
                    values = [item[field] for item in rows]
                    workload_rows[-1].update({f"{field}_mean": sum(values) / k,
                                              f"{field}_min": min(values), f"{field}_max": max(values),
                                              f"{field}_cv": _cv(values)})
            for site in row["sites"].split(","):
                site_mapping_rows.append({"K": k, "site": site, "group": row["group"],
                                          "atomic_unit": next(",".join(sorted(unit)) for unit in units if site in unit)})
            for sortie_id in row["transport_sorties"].split(","):
                if sortie_id: mapping_rows.append({"K": k, "group": row["group"], "task_type": "transport", "task_id": sortie_id})
            for sortie_id in row["relay_sorties"].split(","):
                if sortie_id: mapping_rows.append({"K": k, "group": row["group"], "task_type": "relay", "task_id": sortie_id})
        intervals_by_key = interval_data
        totals = {key: sum(row[key] for row in rows) for key in RESOURCE_KEYS}
        for key in RESOURCE_KEYS:
            global_peak = _peak([(a, b) for a, b, _ in intervals_by_key[key]])
            inventory_rows.append({"K": k, "resource": key, "inventory": inventory[key], "group_demands": ",".join(str(row[key]) for row in rows),
                                  "total_demand": totals[key], "global_peak": global_peak,
                                  "partition_overhead": totals[key] - global_peak,
                                  "inventory_deficit": max(0, totals[key] - inventory[key]),
                                  "inventory_surplus": max(0, inventory[key] - totals[key]),
                                  "deficit_reason": "冻结各组任务时序后，资源不得跨组复用，各组占用峰值之和超过总库存" if totals[key] > inventory[key] else "无库存缺口",
                                  "inventory_scope": "全局总库存；未指定初始分组配给，不将同一库存重复计为各组余量"})
            for start, end, task in intervals_by_key[key]:
                owners = [row["group"] for row in rows if task in
                          (row["transport_sorties"] + "," + row["relay_sorties"]).split(",")]
                for task_group in owners:
                    interval_rows.append({"K": k, "group": task_group, "resource": key, "task_id": task,
                                          "start_s": start, "end_s": end, "half_open": True,
                                          "shared_across_groups": len(owners) > 1})
    _write_csv(output_dir / "problem4_configuration.csv", template_rows)
    _write_csv(output_dir / "inventory_comparison.csv", inventory_rows)
    _write_csv(output_dir / "workload_comparison.csv", workload_rows)
    _write_csv(output_dir / "task_mapping_audit.csv", mapping_rows)
    _write_csv(output_dir / "site_mapping_audit.csv", site_mapping_rows)
    _write_csv(output_dir / "resource_interval_audit.csv", interval_rows)
    _write_submission_workbook(output_dir / "problem4_submission.xlsx", template_rows)
    (output_dir / "assumptions.json").write_text(json.dumps({
        "q3_fixed_schedule": True,
        "partition_selection": "lexicographic: total inventory deficit, total resource demand, maximum CV of transport time/relay service time/sortie count/box count",
        "relay_partition_rule": "merge sites sharing a transport sortie or a relay mission transitively; preserve every Q3 mission and its full communication mapping exactly once; no duplication, splitting, re-timing or cross-group resource sharing",
        "partition_policy": PARTITION_POLICY,
        "inventory_comparison": "resource counts are kept separately by type; aggregate deficit and scale are unweighted modeling objectives, not procurement costs",
        "resource_reassignment": "task start/end times remain fixed; identical units may be recolored within each group; no resource is reused across groups",
        "optimality_scope": "complete partition enumeration for this frozen Q3 schedule and the stated objective only",
        "workload_statistics": "population CV (standard deviation divided by mean); undefined at zero mean",
        "resource_intervals": "half-open; batteries and relay energy components held through full recharge; relay airframe includes audited turnaround",
        "mass_by_group": "summed from Q3 box_delivery_audit.csv with unique-box and route/count checks when available; omitted if no cargo audit exists",
        "partition_search_limit": None,
        "inventory_source": "transport and relay source workbooks"
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    (output_dir / "result.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    return result


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    print(json.dumps(run(args.input, args.output), ensure_ascii=False, indent=2))
