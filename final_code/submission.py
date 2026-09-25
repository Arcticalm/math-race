"""Classify all results using the official template, without mixing Q2 and Q3."""
from __future__ import annotations

import csv
from pathlib import Path
from openpyxl import load_workbook

from final_code.problem1.solver import ROOT

QUESTION_SHEETS = {
    'q1': ('Q1_单点组批',),
    'q2': ('Q2_运输架次', 'Q2_逐箱交付'),
    'q3': ('Q2_运输架次', 'Q2_逐箱交付', 'Q3_中继架次', 'Q3_通信保障'),
    'q4': ('Q4_分区配置',),
}


def _headers(sheet):
    values = [cell.value for cell in sheet[1]]
    while values and values[-1] is None:
        values.pop()
    return values


def _rows(sheet, width):
    return [tuple(row) for row in sheet.iter_rows(min_row=2, max_col=width, values_only=True)
            if any(value is not None for value in row)]


def _fill(sheet, rows):
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions


def organize_workbooks(output: Path, template: Path | None = None) -> dict:
    """Validate all inputs before replacing the four question workbooks."""
    template = template or ROOT / 'docs/结果提交模板.xlsx'
    combined = load_workbook(template)
    headers = {name: _headers(combined[name]) for name in combined.sheetnames}
    data = {}
    for question, names in QUESTION_SHEETS.items():
        path = output / question / f'problem{question[1:]}_submission.xlsx'
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for name in names:
                if name not in workbook or _headers(workbook[name]) != headers[name]:
                    raise ValueError(f'Template header mismatch: {question}/{name}')
                data[question, name] = _rows(workbook[name], len(headers[name]))
        finally:
            workbook.close()
    counts = {}
    for question, names in QUESTION_SHEETS.items():
        workbook = load_workbook(template)
        for name in list(workbook.sheetnames):
            if name not in names:
                del workbook[name]
        tables = output / question / 'tables'
        tables.mkdir(exist_ok=True)
        counts[question] = {}
        for name in names:
            rows = data[question, name]
            _fill(workbook[name], rows)
            label = name.replace('Q2_', 'Q3_') if question == 'q3' else name
            with (tables / f'{label}.csv').open('w', encoding='utf-8-sig', newline='') as stream:
                writer = csv.writer(stream, lineterminator='\n')
                writer.writerow(headers[name])
                writer.writerows(rows)
            counts[question][label] = len(rows)
        workbook.save(output / question / f'problem{question[1:]}_submission.xlsx')
        workbook.close()
    for name in combined.sheetnames:
        question = name[:2].lower()
        _fill(combined[name], data[question, name])
    # The template has no dedicated transport sheets for Q3. Append these two
    # supplements so the independently selected Q2 solution remains unchanged.
    for name in ('Q2_运输架次', 'Q2_逐箱交付'):
        sheet = combined.copy_worksheet(combined[name])
        sheet.title = name.replace('Q2_', 'Q3_')
        _fill(sheet, data['q3', name])
    combined.save(output / '结果提交汇总.xlsx')
    combined.close()
    return counts
