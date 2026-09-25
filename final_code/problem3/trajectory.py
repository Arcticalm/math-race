"""Transport trajectories, deadline checks and Question 3 output helpers."""
from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import load_workbook

from final_code.problem1.solver import ROOT, Node, load_aircraft
from final_code.problem2.transport import RouteEvaluator, _hard_deadlines

REFERENCE_RELAY_POSITIONS = {
    "P1": (109.19944444, 23.05750000),
    "P2": (109.28027778, 23.03333333),
    "P3": (109.19444444, 23.06333333),
}


@dataclass(frozen=True)
class TrackPhase:
    sortie: str
    name: str
    start_s: float
    end_s: float
    start_node: Node
    end_node: Node
    start_altitude_m: float
    end_altitude_m: float


def _node_at(first: Node, second: Node, fraction: float, code: str) -> Node:
    return Node(code, first.longitude + fraction * (second.longitude - first.longitude),
                first.latitude + fraction * (second.latitude - first.latitude), 0.0)


def build_trajectory(sorties, reserve: float = 0.2, energy_scale: float = 1.0) -> list[TrackPhase]:
    aircrafts = load_aircraft()
    base_evaluator = RouteEvaluator(reserve, energy_scale)
    try:
        base, sites = base_evaluator.base, base_evaluator.sites
        phases = []
        for sortie in sorties:
            aircraft = aircrafts[sortie.aircraft]
            current_node = base
            current_altitude = base.elevation_m
            current_time = sortie.takeoff_s
            for sequence, leg in enumerate(sortie.legs or [], 1):
                destination = base if leg["to"] == "O01" else sites[leg["to"]]
                destination_altitude = destination.elevation_m + (30 if destination.code != "O01" else 0)
                cruise_altitude = leg["cruise_altitude_m"]
                distance = leg["distance_m"]
                climb_s = max(0.0, cruise_altitude - current_altitude) / aircraft.climb_speed_mps
                cruise_s = distance / aircraft.cruise_speed_mps
                descent_s = max(0.0, cruise_altitude - destination_altitude) / aircraft.descent_speed_mps
                phases.append(TrackPhase(sortie.code, f"leg{sequence}-climb", current_time,
                    current_time + climb_s, current_node, current_node, current_altitude, cruise_altitude))
                current_time += climb_s
                phases.append(TrackPhase(sortie.code, f"leg{sequence}-cruise", current_time,
                    current_time + cruise_s, current_node, destination, cruise_altitude, cruise_altitude))
                current_time += cruise_s
                phases.append(TrackPhase(sortie.code, f"leg{sequence}-descent", current_time,
                    current_time + descent_s, destination, destination, cruise_altitude, destination_altitude))
                current_time += descent_s
                current_node, current_altitude = destination, destination_altitude
                if destination.code != "O01":
                    count = sum(box.site == destination.code for box in sortie.boxes)
                    handoff = aircraft.handoff_base_s + aircraft.handoff_each_s * count
                    phases.append(TrackPhase(sortie.code, f"{destination.code}-handoff", current_time,
                        current_time + handoff, destination, destination, destination_altitude, destination_altitude))
                    current_time += handoff
            if abs(current_time - sortie.return_s) > 1e-6:
                raise ValueError(f"Trajectory timeline mismatch for {sortie.code}: {current_time} != {sortie.return_s}")
        return phases
    finally:
        base_evaluator.close()


def _position(phase: TrackPhase, time_s: float) -> tuple[Node, float]:
    duration = phase.end_s - phase.start_s
    fraction = 0.0 if duration <= 0 else min(1.0, max(0.0, (time_s - phase.start_s) / duration))
    node = _node_at(phase.start_node, phase.end_node, fraction, phase.name)
    altitude = phase.start_altitude_m + fraction * (phase.end_altitude_m - phase.start_altitude_m)
    return node, altitude


def _hard_deadline_audit(sorties) -> dict:
    checks = []
    for sortie in sorties:
        deliveries = sortie.deliveries or {}
        for box in sortie.boxes:
            delivered = deliveries.get(box.code)
            for kind, deadline_s in _hard_deadlines(box):
                checks.append({
                    "sortie": sortie.code, "box": box.code, "kind": kind,
                    "delivery_s": delivered, "deadline_s": deadline_s,
                    "passed": delivered is not None and delivered <= deadline_s + 1e-7,
                })
    return {
        "check_count": len(checks),
        "passed_count": sum(item["passed"] for item in checks),
        "failed_count": sum(not item["passed"] for item in checks),
        "passed": all(item["passed"] for item in checks),
        "checks": checks,
    }


def _write_csv(path: Path, rows: list[dict]) -> None:
    if not rows:
        path.write_text("", encoding="utf-8-sig")
        return
    fieldnames = list(dict.fromkeys(key for row in rows for key in row))
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fieldnames, extrasaction="ignore", lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def _write_q3_submission(output_dir: Path, sorties, selected, communication_rows) -> None:
    workbook = load_workbook(ROOT / "docs" / "结果提交模板.xlsx")
    transport_sheet = workbook["Q2_运输架次"]
    if transport_sheet.max_row > 1:
        transport_sheet.delete_rows(2, transport_sheet.max_row - 1)
    for sortie in sorties:
        transport_sheet.append([
            sortie.code, sortie.drone, sortie.aircraft, sortie.battery,
            round(sortie.prep_start_s, 3), "→".join(["O01", *sortie.route, "O01"]),
            round(sortie.return_s, 3), round(sortie.energy_kwh, 6),
        ])
    delivery_sheet = workbook["Q2_逐箱交付"]
    if delivery_sheet.max_row > 1:
        delivery_sheet.delete_rows(2, delivery_sheet.max_row - 1)
    for sortie in sorties:
        for box in sortie.boxes:
            delivery_sheet.append([
                box.code, sortie.code, box.site,
                round((sortie.deliveries or {})[box.code], 3),
            ])
    relay_sheet = workbook["Q3_中继架次"]
    if relay_sheet.max_row > 1:
        relay_sheet.delete_rows(2, relay_sheet.max_row - 1)
    for item in selected:
        relay_sheet.append([
            item.get("relay_sortie", ""), item.get("relay_drone", ""),
            item.get("energy_component", ""), round(item["preparation_start_s"], 3),
            item["longitude"], item["latitude"], item["hover_altitude_m"],
            round(item["service_start_s"], 3), round(item["service_end_s"], 3),
            round(item["return_o01_s"], 3), round(item["energy_kwh"], 6),
        ])
    communication_sheet = workbook["Q3_通信保障"]
    if communication_sheet.max_row > 1:
        communication_sheet.delete_rows(2, communication_sheet.max_row - 1)
    for row in communication_rows:
        communication_sheet.append([row[key] for key in (
            "运输架次编号", "通信阶段", "开始时刻（s）", "结束时刻（s）",
            "保障方式", "中继架次编号")])
    workbook.save(output_dir / "problem3_submission.xlsx")


def _plot_problem3_overview(output_path: Path, phases: list[TrackPhase], selected: list[dict]) -> None:
    figure, axis = plt.subplots(figsize=(10, 7), layout="constrained")
    seen = set()
    for phase in phases:
        if phase.sortie in seen or phase.start_node.code == phase.end_node.code:
            continue
        seen.add(phase.sortie)
        axis.plot([phase.start_node.longitude, phase.end_node.longitude],
                  [phase.start_node.latitude, phase.end_node.latitude],
                  linewidth=0.8, alpha=0.45)
    if selected:
        axis.scatter([item["longitude"] for item in selected],
                     [item["latitude"] for item in selected],
                     marker="^", s=55, color="#c0392b", label="Relay hover point")
    axis.set_xlabel("Longitude")
    axis.set_ylabel("Latitude")
    axis.set_title("Problem 3 transport routes and relay schedule")
    if selected:
        axis.legend()
    figure.savefig(output_path, dpi=160)
    plt.close(figure)

