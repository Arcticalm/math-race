from __future__ import annotations

import argparse
import csv
import json
import math
from dataclasses import asdict, dataclass
from pathlib import Path

import rasterio
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import load_workbook

from final_code.problem1.solver import Node, load_aircraft
from final_code.problem2.transport import (
    RouteEvaluator,
    _hard_deadlines,
    _validate,
    load_resources,
)
from final_code.problem3.physics import (
    LinkEvaluator,
    sampled_flight_leg,
    estimate_relay_mission,
    link_limits,
    load_link_parameters,
    load_relay_parameters,
)

ROOT = Path(__file__).resolve().parents[2]
OUTPUT_DEFAULT = ROOT / "outputs" / "problem3"
REFERENCE_RELAY_POSITIONS = {
    "P1": (109.19944444, 23.05750000),
    "P2": (109.28027778, 23.03333333),
    "P3": (109.19444444, 23.06333333),
}


@dataclass(frozen=True)
class TrackPhase:
    sortie: str
    name: str
    start_s: float
    end_s: float
    start_node: Node
    end_node: Node
    start_altitude_m: float
    end_altitude_m: float


def _node_at(first: Node, second: Node, fraction: float, code: str) -> Node:
    return Node(code, first.longitude + fraction * (second.longitude - first.longitude),
                first.latitude + fraction * (second.latitude - first.latitude), 0.0)


def build_trajectory(sorties, reserve: float = 0.2, energy_scale: float = 1.0) -> list[TrackPhase]:
    aircrafts = load_aircraft()
    base_evaluator = RouteEvaluator(reserve, energy_scale)
    try:
        base, sites = base_evaluator.base, base_evaluator.sites
        phases = []
        for sortie in sorties:
            aircraft = aircrafts[sortie.aircraft]
            current_node = base
            current_altitude = base.elevation_m
            current_time = sortie.takeoff_s
            for sequence, leg in enumerate(sortie.legs or [], 1):
                destination = base if leg["to"] == "O01" else sites[leg["to"]]
                destination_altitude = destination.elevation_m + (30 if destination.code != "O01" else 0)
                cruise_altitude = leg["cruise_altitude_m"]
                distance = leg["distance_m"]
                climb_s = max(0.0, cruise_altitude - current_altitude) / aircraft.climb_speed_mps
                cruise_s = distance / aircraft.cruise_speed_mps
                descent_s = max(0.0, cruise_altitude - destination_altitude) / aircraft.descent_speed_mps
                phases.append(TrackPhase(sortie.code, f"leg{sequence}-climb", current_time,
                    current_time + climb_s, current_node, current_node, current_altitude, cruise_altitude))
                current_time += climb_s
                phases.append(TrackPhase(sortie.code, f"leg{sequence}-cruise", current_time,
                    current_time + cruise_s, current_node, destination, cruise_altitude, cruise_altitude))
                current_time += cruise_s
                phases.append(TrackPhase(sortie.code, f"leg{sequence}-descent", current_time,
                    current_time + descent_s, destination, destination, cruise_altitude, destination_altitude))
                current_time += descent_s
                current_node, current_altitude = destination, destination_altitude
                if destination.code != "O01":
                    count = sum(box.site == destination.code for box in sortie.boxes)
                    handoff = aircraft.handoff_base_s + aircraft.handoff_each_s * count
                    phases.append(TrackPhase(sortie.code, f"{destination.code}-handoff", current_time,
                        current_time + handoff, destination, destination, destination_altitude, destination_altitude))
                    current_time += handoff
            if abs(current_time - sortie.return_s) > 1e-6:
                raise ValueError(f"Trajectory timeline mismatch for {sortie.code}: {current_time} != {sortie.return_s}")
        return phases
    finally:
        base_evaluator.close()


def _position(phase: TrackPhase, time_s: float) -> tuple[Node, float]:
    duration = phase.end_s - phase.start_s
    fraction = 0.0 if duration <= 0 else min(1.0, max(0.0, (time_s - phase.start_s) / duration))
    node = _node_at(phase.start_node, phase.end_node, fraction, phase.name)
    altitude = phase.start_altitude_m + fraction * (phase.end_altitude_m - phase.start_altitude_m)
    return node, altitude


def screen_direct_links(sorties, phases: list[TrackPhase], step_s: float = 10.0) -> tuple[list[dict], list[dict]]:
    params = load_link_parameters()
    evaluator = LinkEvaluator(params)
    base, _ = evaluator.base, evaluator.sites
    gateway = Node("G01", base.longitude, base.latitude, base.elevation_m + evaluator.params.gateway_agl_m)
    threshold = link_limits(params)["transport_gateway_db"]
    samples = []
    try:
        for phase in phases:
            if phase.end_s <= phase.start_s:
                continue
            sample_count = max(1, math.ceil((phase.end_s - phase.start_s) / step_s))
            for index in range(sample_count + 1):
                time_s = phase.start_s + (phase.end_s - phase.start_s) * index / sample_count
                if index == sample_count and phase.end_s == phase.start_s:
                    continue
                position, altitude = _position(phase, time_s)
                link = evaluator.evaluate(position, altitude, gateway, gateway.elevation_m, threshold)
                samples.append({
                    "sortie": phase.sortie, "phase": phase.name, "time_s": time_s,
                    "longitude": position.longitude, "latitude": position.latitude,
                    "altitude_m": altitude, **asdict(link),
                })
    finally:
        evaluator.close()
    intervals = _sample_intervals(samples)
    return samples, intervals


def _sample_intervals(samples: list[dict]) -> list[dict]:
    """Cover every sample cell; either failed endpoint requests relay coverage."""
    groups = {}
    for sample in samples:
        groups.setdefault((sample["sortie"], sample["phase"]), []).append(sample)
    intervals = []
    for (sortie, phase), points in groups.items():
        points.sort(key=lambda item: item["time_s"])
        for first, second in zip(points, points[1:]):
            if second["time_s"] <= first["time_s"]:
                continue
            available = first["available"] and second["available"]
            reason = first["reason"] or second["reason"]
            if (intervals and intervals[-1]["sortie"] == sortie
                    and intervals[-1]["phase"] == phase
                    and intervals[-1]["direct_available"] == available
                    and intervals[-1]["reason"] == reason):
                intervals[-1]["end_s"] = second["time_s"]
            else:
                intervals.append({"sortie": sortie, "phase": phase,
                                  "start_s": first["time_s"], "end_s": second["time_s"],
                                  "direct_available": available, "reason": reason})
    return sorted(intervals, key=lambda item: (item["start_s"], item["sortie"]))


def _merge_gap_intervals(intervals: list[dict]) -> list[dict]:
    merged = []
    by_sortie = {}
    for interval in intervals:
        if interval["direct_available"] or interval["end_s"] <= interval["start_s"]:
            continue
        by_sortie.setdefault(interval["sortie"], []).append(interval)
    for sortie, values in by_sortie.items():
        values.sort(key=lambda item: item["start_s"])
        current = None
        for value in values:
            if current is None or value["start_s"] > current["end_s"] + 1e-6:
                if current is not None:
                    merged.append(current)
                current = {
                    "sortie": value["sortie"], "start_s": value["start_s"],
                    "end_s": value["end_s"], "direct_available": False,
                    "reason": value.get("reason", ""),
                }
            else:
                current["end_s"] = max(current["end_s"], value["end_s"])
        if current is not None:
            merged.append(current)
    return sorted(merged, key=lambda item: (item["start_s"], item["sortie"]))


def _assign_gap_ids(intervals: list[dict], gaps: list[dict]) -> None:
    for gap_id, gap in enumerate(gaps, 1):
        gap["gap_id"] = gap_id
    for interval in intervals:
        interval.pop("gap_id", None)
        if interval["direct_available"]:
            continue
        matches = [gap["gap_id"] for gap in gaps
                   if gap["sortie"] == interval["sortie"]
                   and interval["start_s"] < gap["end_s"] - 1e-7
                   and interval["end_s"] > gap["start_s"] + 1e-7]
        if matches:
            interval["gap_id"] = matches[0]


def _hard_deadline_audit(sorties) -> dict:
    checks = []
    for sortie in sorties:
        deliveries = sortie.deliveries or {}
        for box in sortie.boxes:
            delivered = deliveries.get(box.code)
            for kind, deadline_s in _hard_deadlines(box):
                checks.append({
                    "sortie": sortie.code, "box": box.code, "kind": kind,
                    "delivery_s": delivered, "deadline_s": deadline_s,
                    "passed": delivered is not None and delivered <= deadline_s + 1e-7,
                })
    return {
        "check_count": len(checks),
        "passed_count": sum(item["passed"] for item in checks),
        "failed_count": sum(not item["passed"] for item in checks),
        "passed": all(item["passed"] for item in checks),
        "checks": checks,
    }


def _shift_sortie(sortie, delay_s: float) -> None:
    sortie.prep_start_s += delay_s
    sortie.prep_end_s += delay_s
    sortie.loading_end_s += delay_s
    sortie.takeoff_s += delay_s
    sortie.return_s += delay_s
    sortie.deliveries = {code: time_s + delay_s for code, time_s in (sortie.deliveries or {}).items()}


def search_relay_candidates(sorties, phases: list[TrackPhase], gaps: list[dict],
                            sample_step_s: float = 30.0,
                            candidate_window_s: float = 20000.0,
                            checkpoints: list[dict] | None = None,
                            allow_early_departure: bool = False) -> list[dict]:
    params = load_link_parameters()
    relay = load_relay_parameters()
    evaluator = LinkEvaluator(params)
    base = evaluator.base
    gateway = Node("G01", base.longitude, base.latitude, base.elevation_m + evaluator.params.gateway_agl_m)
    limits = link_limits(params)
    time_step_s = min(30.0, sample_step_s)
    offsets = (-0.01, 0.0, 0.01)
    results = []
    phases_by_sortie: dict[str, list[TrackPhase]] = {}
    sorties_by_code = {sortie.code: sortie for sortie in sorties}
    for phase in phases:
        phases_by_sortie.setdefault(phase.sortie, []).append(phase)
    try:
        gap_groups = []
        for gap_id, gap in enumerate(sorted(gaps, key=lambda item: (item["start_s"], item["sortie"])), 1):
            gap["gap_id"] = gap_id
            if (not gap_groups
                    or gap["start_s"] - gap_groups[-1][0]["start_s"] > candidate_window_s + 1e-7):
                gap_groups.append([gap])
            else:
                gap_groups[-1].append(gap)

        for group_index, group in enumerate(gap_groups, 1):
            demand_by_gap = {}
            target_nodes = []
            for gap in group:
                sortie_phases = phases_by_sortie[gap["sortie"]]
                demand_by_gap[gap["gap_id"]] = []
                sample_count = max(1, math.ceil((gap["end_s"] - gap["start_s"]) / time_step_s))
                for sample_index in range(sample_count + 1):
                    time_s = gap["start_s"] + (gap["end_s"] - gap["start_s"]) * sample_index / sample_count
                    phase = next((item for item in sortie_phases
                                  if item.start_s - 1e-7 <= time_s <= item.end_s + 1e-7), None)
                    if phase is None:
                        continue
                    transport_node, transport_altitude = _position(phase, time_s)
                    demand_by_gap[gap["gap_id"]].append((time_s, transport_node, transport_altitude))
                # Include every original trajectory checkpoint and all phase boundaries.
                for phase in sortie_phases:
                    for time_s in (phase.start_s, phase.end_s):
                        if gap["start_s"] - 1e-7 <= time_s <= gap["end_s"] + 1e-7:
                            node, altitude = _position(phase, time_s)
                            demand_by_gap[gap["gap_id"]].append((time_s, node, altitude))
                for point in checkpoints or []:
                    if (point["sortie"] == gap["sortie"]
                            and gap["start_s"] - 1e-7 <= point["time_s"] <= gap["end_s"] + 1e-7):
                        demand_by_gap[gap["gap_id"]].append((point["time_s"],
                            Node("checkpoint", point["longitude"], point["latitude"], 0), point["altitude_m"]))
                for fraction in (0.0, 0.25, 0.5, 0.75, 1.0):
                    time_s = gap["start_s"] + (gap["end_s"] - gap["start_s"]) * fraction
                    phase = next(item for item in sortie_phases
                                 if item.start_s - 1e-7 <= time_s <= item.end_s + 1e-7)
                    target_nodes.append(_position(phase, time_s)[0])
                unique_demand = {}
                for time_s, node, altitude in demand_by_gap[gap["gap_id"]]:
                    key = (round(time_s, 7), round(node.longitude, 10),
                           round(node.latitude, 10), round(altitude, 5))
                    unique_demand.setdefault(key, (time_s, node, altitude))
                demand_by_gap[gap["gap_id"]] = list(unique_demand.values())
            locations = {(round(node.longitude + lon_offset, 7), round(node.latitude + lat_offset, 7))
                         for node in target_nodes for lon_offset in offsets for lat_offset in offsets}
            # The reference solver's P1/P2/P3 points are retained as explicit
            # candidates. They are service-area relay points derived from the
            # same DEM and communication workbooks, not fabricated coverage.
            locations.update(REFERENCE_RELAY_POSITIONS.values())
            candidate_rows = []
            for longitude, latitude in sorted(locations):
                row, col = rasterio.transform.rowcol(evaluator.dem.transform, longitude, latitude)
                if not (0 <= row < evaluator.dem.height and 0 <= col < evaluator.dem.width):
                    continue
                ground_m = float(evaluator.dem_values[row, col])
                if not math.isfinite(ground_m) or math.isclose(ground_m, -32767.0):
                    continue
                hover = Node("H", longitude, latitude, ground_m)
                try:
                    outbound_max_ground, outbound_distance = sampled_flight_leg(evaluator, base, hover)
                    return_max_ground, return_distance = sampled_flight_leg(evaluator, hover, base)
                except ValueError:
                    continue
                for agl_m in (50.0, 150.0, 250.0, relay.maximum_agl_m):
                    if agl_m > relay.maximum_agl_m:
                        continue
                    hover = Node("H", longitude, latitude, ground_m)
                    hover_altitude = ground_m + agl_m
                    backhaul = evaluator.evaluate(hover, hover_altitude, gateway,
                                                   gateway.elevation_m, limits["relay_gateway_db"])
                    if not backhaul.available:
                        continue
                    covered = []
                    for gap in group:
                        access_ok = all(
                            evaluator.evaluate(transport_node, transport_altitude, hover,
                                               hover_altitude, limits["transport_relay_db"]).available
                            for _, transport_node, transport_altitude in demand_by_gap[gap["gap_id"]]
                        )
                        if demand_by_gap[gap["gap_id"]] and access_ok:
                            covered.append(gap)
                    if not covered:
                        continue
                    from final_code.problem3.timeline import overlapping_windows
                    for window_subset in overlapping_windows(covered):
                        service_start_actual = min(item["start_s"] for item in window_subset)
                        service_end = max(item["end_s"] for item in window_subset)
                        service_s = service_end - service_start_actual
                        mission = estimate_relay_mission(
                            service_s, outbound_distance, outbound_max_ground, hover_altitude,
                            return_distance, return_max_ground, base.elevation_m, relay,
                        )
                        prepare_start = service_start_actual - relay.preparation_s - mission.outbound_flight_s - relay.link_setup_s
                        if (prepare_start < -1e-7 and not allow_early_departure) or not mission.feasible_energy:
                            continue
                        subset_ids = [item["gap_id"] for item in window_subset]
                        candidate_rows.append({
                            "covered_gap_ids": subset_ids,
                            "spatial_gap_ids": [item["gap_id"] for item in covered],
                            "covered_sorties": ",".join(sorted({item["sortie"] for item in window_subset})),
                            "service_start_s": service_start_actual, "service_end_s": service_end,
                            "preparation_start_s": prepare_start,
                            "longitude": longitude, "latitude": latitude,
                            "ground_dsm_m": ground_m, "hover_altitude_m": hover_altitude,
                            "agl_m": agl_m, "energy_kwh": mission.energy_kwh,
                            "return_soc_percent": mission.return_soc_percent,
                            "return_o01_s": service_end + mission.return_flight_s,
                            "outbound_flight_s": mission.outbound_flight_s,
                            "return_flight_s": mission.return_flight_s,
                        })
            by_coverage = {}
            for candidate in candidate_rows:
                by_coverage.setdefault(tuple(candidate["covered_gap_ids"]), []).append(candidate)
            representatives = []
            for options in by_coverage.values():
                keys = (lambda c: c["energy_kwh"], lambda c: -c["preparation_start_s"],
                        lambda c: c["return_o01_s"])
                chosen = {}
                for key in keys:
                    for c in sorted(options, key=key)[:3]:
                        chosen[(c["longitude"], c["latitude"], c["hover_altitude_m"])] = c
                representatives.extend(chosen.values())
            candidate_rows = sorted(representatives, key=lambda item: (
                -len(item["covered_gap_ids"]), item["energy_kwh"], item["return_o01_s"],
            ))
            results.append({
                "group": group_index, "covered_sorties": ",".join(sorted({item["sortie"] for item in group})),
                "gap_count": len(group), "gap_start_s": min(item["start_s"] for item in group),
                "gap_end_s": max(item["end_s"] for item in group),
                "gap_start_by_id": {str(item["gap_id"]): item["start_s"] for item in group},
                "gaps": [{"gap_id": item["gap_id"], "sortie": item["sortie"],
                          "start_s": item["start_s"], "end_s": item["end_s"]} for item in group],
                "candidate_count": len(candidate_rows), "candidates": candidate_rows,
            })
    finally:
        evaluator.close()
    return results


def _write_csv(path: Path, rows: list[dict]) -> None:
    if not rows:
        path.write_text("", encoding="utf-8-sig")
        return
    fieldnames = list(dict.fromkeys(key for row in rows for key in row))
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fieldnames, extrasaction="ignore", lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def _relay_for_gap(selected: list[dict], gap_id: int) -> dict | None:
    for item in selected:
        if gap_id in item.get("covered_gap_ids", []):
            return item
    return None


def _communication_rows(sorties, intervals, selected) -> list[dict]:
    rows = []
    for interval in intervals:
        if interval["end_s"] <= interval["start_s"]:
            continue
        relay = None
        mode = "直连" if interval["direct_available"] else "中断"
        if not interval["direct_available"] and interval.get("gap_id") is not None:
            relay = _relay_for_gap(selected, int(interval["gap_id"]))
            if relay is not None:
                mode = "中继候选"
        rows.append({
            "运输架次编号": interval["sortie"],
            "通信阶段": interval.get("phase", "空中通信"),
            "开始时刻（s）": interval["start_s"],
            "结束时刻（s）": interval["end_s"],
            "保障方式": mode,
            "中继架次编号": relay.get("relay_sortie", "") if relay else "",
            "诊断说明": interval.get("reason", "") if mode == "中断" else "",
        })
    return rows


def _write_q3_submission(output_dir: Path, sorties, intervals, selected, communication_rows=None) -> None:
    workbook = load_workbook(ROOT / "docs" / "结果提交模板.xlsx")
    transport_sheet = workbook["Q2_运输架次"]
    if transport_sheet.max_row > 1:
        transport_sheet.delete_rows(2, transport_sheet.max_row - 1)
    for sortie in sorties:
        transport_sheet.append([
            sortie.code, sortie.drone, sortie.aircraft, sortie.battery,
            round(sortie.prep_start_s, 3), "→".join(["O01", *sortie.route, "O01"]),
            round(sortie.return_s, 3), round(sortie.energy_kwh, 6),
        ])
    delivery_sheet = workbook["Q2_逐箱交付"]
    if delivery_sheet.max_row > 1:
        delivery_sheet.delete_rows(2, delivery_sheet.max_row - 1)
    for sortie in sorties:
        for box in sortie.boxes:
            delivery_sheet.append([
                box.code, sortie.code, box.site,
                round((sortie.deliveries or {})[box.code], 3),
            ])
    relay_sheet = workbook["Q3_中继架次"]
    if relay_sheet.max_row > 1:
        relay_sheet.delete_rows(2, relay_sheet.max_row - 1)
    for item in selected:
        relay_sheet.append([
            item.get("relay_sortie", ""), item.get("relay_drone", ""),
            item.get("energy_component", ""), round(item["preparation_start_s"], 3),
            item["longitude"], item["latitude"], item["hover_altitude_m"],
            round(item["service_start_s"], 3), round(item["service_end_s"], 3),
            round(item["return_o01_s"], 3), round(item["energy_kwh"], 6),
        ])
    communication_sheet = workbook["Q3_通信保障"]
    if communication_sheet.max_row > 1:
        communication_sheet.delete_rows(2, communication_sheet.max_row - 1)
    for row in (communication_rows if communication_rows is not None else _communication_rows(sorties, intervals, selected)):
        communication_sheet.append([row[key] for key in (
            "运输架次编号", "通信阶段", "开始时刻（s）", "结束时刻（s）",
            "保障方式", "中继架次编号")])
    workbook.save(output_dir / "problem3_submission.xlsx")


def _plot_problem3_overview(output_path: Path, phases: list[TrackPhase], selected: list[dict]) -> None:
    figure, axis = plt.subplots(figsize=(10, 7), layout="constrained")
    seen = set()
    for phase in phases:
        if phase.sortie in seen or phase.start_node.code == phase.end_node.code:
            continue
        seen.add(phase.sortie)
        axis.plot([phase.start_node.longitude, phase.end_node.longitude],
                  [phase.start_node.latitude, phase.end_node.latitude],
                  linewidth=0.8, alpha=0.45)
    if selected:
        axis.scatter([item["longitude"] for item in selected],
                     [item["latitude"] for item in selected],
                     marker="^", s=55, color="#c0392b", label="Relay hover point")
    axis.set_xlabel("Longitude")
    axis.set_ylabel("Latitude")
    axis.set_title("Problem 3 transport routes and relay schedule")
    if selected:
        axis.legend()
    figure.savefig(output_path, dpi=160)
    plt.close(figure)


def _package_program(output_dir: Path) -> None:
    from final_code.package import write_program_bundle

    write_program_bundle(output_dir / "problem3_program.zip")


def _node_elevation_audit(evaluator: LinkEvaluator) -> list[dict]:
    nodes = [evaluator.base, *evaluator.sites.values()]
    audit = []
    for node in nodes:
        row, col = rasterio.transform.rowcol(evaluator.dem.transform, node.longitude, node.latitude)
        valid = 0 <= row < evaluator.dem.height and 0 <= col < evaluator.dem.width
        dsm = float(evaluator.dem_values[row, col]) if valid else math.nan
        valid = valid and math.isfinite(dsm) and not math.isclose(dsm, -32767.0)
        audit.append({
            "node": node.code, "node_table_ground_m": node.elevation_m,
            "dsm_pixel_m": dsm if valid else "", "dsm_minus_node_m": dsm - node.elevation_m if valid else "",
            "valid": valid,
        })
    return audit


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def _charge_duration(full_charge_s: float, soc: float) -> float:
    if soc < 0.9:
        return full_charge_s * (0.65 * (0.9 - soc) / 0.9 + 0.35)
    return full_charge_s * 0.35 * (1 - soc) / 0.1


def _prune_dominated_candidates(candidates: list[dict]) -> list[dict]:
    """Drop same-hover missions with worse resources and energy."""
    relay = load_relay_parameters()
    groups = {}
    for candidate in candidates:
        geometry = (candidate.get("longitude"), candidate.get("latitude"),
                    candidate.get("hover_altitude_m"))
        if any(value is None for value in geometry):
            geometry = ("unknown", id(candidate))
        key = (tuple(sorted(candidate["covered_gap_ids"])),
               candidate["service_start_s"], candidate["service_end_s"],
               *geometry)
        groups.setdefault(key, []).append(candidate)

    retained = []
    tolerance = 1e-7
    for options in groups.values():
        for candidate in options:
            candidate_charge_end = candidate["return_o01_s"] + _charge_duration(
                relay.full_charge_s, candidate["return_soc_percent"] / 100)
            dominated = False
            for other in options:
                if other is candidate:
                    continue
                other_charge_end = other["return_o01_s"] + _charge_duration(
                    relay.full_charge_s, other["return_soc_percent"] / 100)
                no_worse = (
                    other["preparation_start_s"] >= candidate["preparation_start_s"] - tolerance
                    and other["return_o01_s"] <= candidate["return_o01_s"] + tolerance
                    and other_charge_end <= candidate_charge_end + tolerance
                    and other["energy_kwh"] <= candidate["energy_kwh"] + tolerance
                )
                strictly_better = (
                    other["preparation_start_s"] > candidate["preparation_start_s"] + tolerance
                    or other["return_o01_s"] < candidate["return_o01_s"] - tolerance
                    or other_charge_end < candidate_charge_end - tolerance
                    or other["energy_kwh"] < candidate["energy_kwh"] - tolerance
                )
                if no_worse and strictly_better:
                    dominated = True
                    break
            if not dominated:
                retained.append(candidate)
    return retained


def _maximal_overlap_cliques(starts, ends, capacity: int) -> list[tuple[int, ...]]:
    """Return maximal overloaded cliques for half-open resource intervals."""
    masks = set()
    for time_s in sorted(set(starts)):
        active = 0
        for index, (start, end) in enumerate(zip(starts, ends)):
            if start <= time_s + 1e-9 and end > time_s + 1e-7:
                active |= 1 << index
        if active.bit_count() > capacity:
            masks.add(active)
    maximal = []
    for mask in sorted(masks, key=int.bit_count, reverse=True):
        if not any(mask & existing == mask for existing in maximal):
            maximal.append(mask)
    return [tuple(index for index in range(len(starts)) if mask & (1 << index)) for mask in maximal]


def select_relay_schedule(candidate_groups: list[dict], gap_count: int,
                          search_limit: int = 50000, objective: str = "energy",
                          max_relay_sorties: int | None = None) -> dict:
    """Select interval missions with exact airframe and component capacities."""
    import numpy as np
    from scipy.optimize import Bounds, LinearConstraint, milp
    from scipy.sparse import lil_matrix

    relay = load_relay_parameters()
    candidates = [{**candidate, "candidate_group": group["group"]}
                  for group in candidate_groups for candidate in group["candidates"]
                  if candidate["preparation_start_s"] >= -1e-7]
    for candidate in candidates:
        if isinstance(candidate["covered_gap_ids"], str):
            candidate["covered_gap_ids"] = json.loads(candidate["covered_gap_ids"])
    candidate_count_before_pruning = len(candidates)
    candidates = _prune_dominated_candidates(candidates)
    impossible = [gap for gap in range(1, gap_count + 1)
                  if not any(gap in c["covered_gap_ids"] for c in candidates)]
    count = len(candidates)
    selected = []
    result = None
    if count:
        starts = np.array([c["preparation_start_s"] for c in candidates])
        drone_ends = np.array([c["return_o01_s"] + relay.turnaround_s for c in candidates])
        battery_ends = np.array([c["return_o01_s"] + _charge_duration(
            relay.full_charge_s, c["return_soc_percent"] / 100) for c in candidates])
        constraints = []
        for gap in range(1, gap_count + 1):
            if gap in impossible:
                continue
            constraints.append(([i for i, c in enumerate(candidates)
                                 if gap in c["covered_gap_ids"]], 1, np.inf))
        if max_relay_sorties is not None:
            constraints.append((list(range(count)), 0, max_relay_sorties))
        # Interval graphs need capacity constraints only on maximal cliques.
        for ends, capacity in ((drone_ends, 2), (battery_ends, relay.battery_count)):
            for active in _maximal_overlap_cliques(starts, ends, capacity):
                constraints.append((active, 0, capacity))
        matrix = lil_matrix((len(constraints), count))
        for row, (indices, _, _) in enumerate(constraints):
            matrix[row, list(indices)] = 1
        energies = np.array([c["energy_kwh"] for c in candidates])
        if objective == "energy":
            costs = energies + 1e-6
        elif objective == "robust":
            relay = load_relay_parameters()
            altitude_penalty = np.array([
                (relay.maximum_agl_m - c["agl_m"]) * 0.001 for c in candidates
            ])
            costs = energies + altitude_penalty
        else:
            costs = np.ones(count) + energies / (float(energies.sum()) + 1)
        result = milp(costs, integrality=np.ones(count), bounds=Bounds(0, 1),
                      constraints=LinearConstraint(matrix.tocsr(),
                          [c[1] for c in constraints], [c[2] for c in constraints]),
                      options={"time_limit": 120.0, "mip_rel_gap": 0.0})
        if result.x is not None:
            selected = [dict(c) for c, x in zip(candidates, result.x) if x > 0.5]
    selected.sort(key=lambda c: c["preparation_start_s"])
    drone_ready = {"R01": 0.0, "R02": 0.0}
    battery_ready = {f"RE-{i:02d}": 0.0 for i in range(1, relay.battery_count + 1)}
    for index, item in enumerate(selected, 1):
        start = item["preparation_start_s"]
        drone = next(code for code, ready in drone_ready.items() if ready <= start + 1e-7)
        battery = next(code for code, ready in battery_ready.items() if ready <= start + 1e-7)
        item.update(relay_sortie=f"R3-{index:03d}", relay_drone=drone, energy_component=battery)
        drone_ready[drone] = item["return_o01_s"] + relay.turnaround_s
        battery_ready[battery] = item["return_o01_s"] + _charge_duration(
            relay.full_charge_s, item["return_soc_percent"] / 100)
        item.update(takeoff_s=start + relay.preparation_s,
                    arrival_s=item["service_start_s"] - relay.link_setup_s,
                    drone_ready_s=drone_ready[drone], charge_complete_s=battery_ready[battery])
    covered = {gap for c in selected for gap in c["covered_gap_ids"]}
    uncovered = sorted(set(range(1, gap_count + 1)) - covered)
    return {"feasible_cover": not uncovered, "selected_count": len(selected),
            "selected": selected, "uncovered_gap_ids": uncovered,
            "impossible_gap_ids": impossible, "best_partial_covered_count": len(covered),
            "candidate_count_before_pruning": candidate_count_before_pruning,
            "candidate_count_after_pruning": len(candidates),
            "drone_ready_times": drone_ready, "energy_component_ready_times": battery_ready,
            "objective": objective,
            "solver_status": result.message if result is not None else "empty or impossible candidate set",
            "mip_gap": float(result.mip_gap) if result is not None and result.x is not None else None,
            "method": "binary interval-cover MILP with airframe and component overlap cliques; interval coloring"}


def run(output_dir: Path = OUTPUT_DEFAULT, sample_step_s: float = 10.0,
        relay_candidate_step_s: float = 10.0,
        q2_output: Path = ROOT / "outputs/problem23/clearance50",
        max_iterations: int = 12, time_limit_s: float = 180.0,
        objective: str = "sorties", max_relay_sorties: int | None = None,
        maximum_delay_s: float = 10800) -> dict:
    from final_code.problem3.timeline import load_problem23_schedule, coordinate_timeline
    from importlib.metadata import version
    import platform
    if (not all(math.isfinite(x) and x > 0 for x in (sample_step_s, relay_candidate_step_s, time_limit_s))
            or max_iterations < 1):
        raise ValueError("Sampling steps, time limit and iteration limit must be positive")
    primary_sorties, boxes, source_metrics = load_problem23_schedule(q2_output)
    drones, batteries = load_resources()
    result = coordinate_timeline(primary_sorties, sample_step_s, relay_candidate_step_s,
                                 max_iterations, time_limit_s, objective, max_relay_sorties,
                                 maximum_delay_s)
    sorties, phases, samples, intervals, gaps, relay_candidates, selected = (
        result[key] for key in ("sorties", "phases", "samples", "intervals", "gaps", "groups", "selected"))
    coordination = result["coordination"]
    candidate_comparison = result["history"]
    flat_candidates = [{"candidate_group": group["group"], **candidate}
                       for group in relay_candidates for candidate in group["candidates"]]
    relay_schedule = result["schedule"]
    transport_evaluator = RouteEvaluator(0.2, 1.0)
    try:
        transport_metrics = _validate(sorties, boxes, drones, batteries, transport_evaluator)
    finally:
        transport_evaluator.close()
    from final_code.problem3.audit import audit_checkpoints, audit_communication, audit_relay_resources
    communication_rows, communication_metrics, link_audit = audit_checkpoints(samples, intervals, selected)
    continuous_rows, continuous_metrics = audit_communication(phases, selected, step_s=30.0, minimum_step_s=1.0)
    relay_metrics = audit_relay_resources(selected)
    hard_deadline_metrics = _hard_deadline_audit(sorties)
    audit_evaluator = LinkEvaluator()
    try:
        elevation_audit = _node_elevation_audit(audit_evaluator)
    finally:
        audit_evaluator.close()
    metrics = {
        "status": "relay candidate/partial schedule diagnostic; not a feasible Q3 solution",
        "transport_feasible": transport_metrics["feasible"],
        "objective": objective, "max_relay_sorties": max_relay_sorties,
        "maximum_transport_delay_s": maximum_delay_s,
        "box_count": len(boxes),
        "sortie_count": len(sorties),
        "direct_link_threshold_db": link_limits(load_link_parameters())["transport_gateway_db"],
        "sample_step_s": sample_step_s,
        "transport_start_coordination": coordination,
        "transport_candidate_comparison": candidate_comparison,
        "sample_count": len(samples),
        "direct_gap_sample_count": sum(not item["available"] for item in samples),
        "direct_gap_interval_count": len(gaps),
        "relay_candidate_step_s": relay_candidate_step_s,
        "relay_candidates_saved_per_group": "energy, departure and return representatives per coverage subset",
        "relay_candidate_summary": [
            {key: gap[key] for key in ("group", "covered_sorties", "gap_count", "gap_start_s", "gap_end_s", "candidate_count")}
            for gap in relay_candidates
        ],
        "gap_groups_with_relay_candidate": sum(item["candidate_count"] > 0 for item in relay_candidates),
        "relay_candidate_feasibility": "all trajectory checkpoints and phase boundaries verified",
        "relay_schedule_feasibility": relay_schedule["method"],
        "relay_schedule": {key: value for key, value in relay_schedule.items() if key != "selected"},
        "DSM_node_elevation_audit_count": len(elevation_audit),
        "continuity_certified": continuous_metrics["certified"],
        "checkpoint_coverage_verified": communication_metrics["certified"],
        "q2_source": source_metrics,
        "software_versions": {"python": platform.python_version(), **{
            package: version(package) for package in ("numpy", "scipy", "rasterio", "ortools", "openpyxl")}},
        "communication_validation": communication_metrics,
        "continuous_communication_validation": continuous_metrics,
        "transport_validation": transport_metrics,
        "hard_deadline_audit": hard_deadline_metrics,
        "relay_validation": relay_metrics,
        "joint_makespan_s": max([s.return_s for s in sorties] + [r["return_o01_s"] for r in relay_schedule["selected"]]),
        "transport_energy_kwh": sum(s.energy_kwh for s in sorties),
        "relay_energy_kwh": sum(r["energy_kwh"] for r in relay_schedule["selected"]),
        "joint_energy_kwh": sum(s.energy_kwh for s in sorties) + sum(r["energy_kwh"] for r in relay_schedule["selected"]),
        "note": "Feasibility requires checkpoint coverage and recursive continuous communication certification under the stated DSM/LOS model.",
    }
    metrics["feasible"] = (transport_metrics["feasible"] and hard_deadline_metrics["passed"]
                           and relay_schedule["feasible_cover"]
                           and communication_metrics["certified"] and continuous_metrics["certified"]
                           and relay_metrics["feasible"])
    if metrics["feasible"]:
        metrics["status"] = "feasible joint schedule with continuous communication certification under the stated DSM/LOS model"
    output_dir.mkdir(parents=True, exist_ok=True)
    _write_csv(output_dir / "link_samples.csv", samples)
    _write_csv(output_dir / "two_hop_link_audit.csv", link_audit)
    _write_csv(output_dir / "continuous_communication_audit.csv", continuous_rows)
    _write_csv(output_dir / "trajectory_phases.csv", [
        {"sortie": p.sortie, "phase": p.name, "start_s": p.start_s, "end_s": p.end_s,
         "start_longitude": p.start_node.longitude, "start_latitude": p.start_node.latitude,
         "end_longitude": p.end_node.longitude, "end_latitude": p.end_node.latitude,
         "start_altitude_m": p.start_altitude_m, "end_altitude_m": p.end_altitude_m}
        for p in phases])
    _write_csv(output_dir / "iteration_history.csv", candidate_comparison)
    _write_csv(output_dir / "direct_link_intervals.csv", intervals)
    _write_csv(output_dir / "relay_candidates.csv", flat_candidates)
    _write_csv(output_dir / "relay_schedule.csv", relay_schedule["selected"])
    _write_csv(output_dir / "node_dsm_elevation_audit.csv", elevation_audit)
    _write_csv(output_dir / "relay_uncovered_gaps.csv", [
        row for row in communication_rows if row["保障方式"] == "中断"])
    _write_csv(output_dir / "gap_id_map.csv", gaps)
    _write_csv(output_dir / "communication_audit.csv",
               communication_rows)
    _write_csv(output_dir / "transport_inherited_audit.csv", [
        {
            "架次编号": sortie.code,
            "无人机编号": sortie.drone,
            "机型编号": sortie.aircraft,
            "电池编号": sortie.battery,
            "准备开始时刻（s）": sortie.prep_start_s,
            "起飞时刻（s）": sortie.takeoff_s,
            "返回O01时刻（s）": sortie.return_s,
            "访问服务区顺序": "→".join(sortie.route),
            "架次能耗（kWh）": sortie.energy_kwh,
            "返航SOC（%）": sortie.battery_soc_return_percent,
            "逐箱交付数": len(sortie.boxes),
            "电池充电完成时刻（s）": sortie.return_s + _charge_duration(
                next(b.full_charge_s for b in batteries if b.code == sortie.battery),
                sortie.battery_soc_return_percent / 100),
        }
        for sortie in sorties
    ])
    _write_csv(output_dir / "relay_resource_audit.csv", [
        {
            "中继架次编号": item.get("relay_sortie", ""),
            "中继无人机编号": item.get("relay_drone", ""),
            "能源组件编号": item.get("energy_component", ""),
            "准备开始时刻（s）": item["preparation_start_s"],
            "建链完成/服务开始（s）": item["service_start_s"],
            "服务结束时刻（s）": item["service_end_s"],
            "返回O01时刻（s）": item["return_o01_s"],
            "机体再次可用（s）": item["return_o01_s"] + load_relay_parameters().turnaround_s,
            "能源组件再次可用（s）": item["return_o01_s"] + _charge_duration(
                load_relay_parameters().full_charge_s, item["return_soc_percent"] / 100),
            "架次能耗（kWh）": item["energy_kwh"],
            "返航SOC（%）": item["return_soc_percent"],
        }
        for item in relay_schedule["selected"]
    ])
    _write_q3_submission(output_dir, sorties, intervals, relay_schedule["selected"], communication_rows)
    _plot_problem3_overview(output_dir / "routes_relays.png", phases, relay_schedule["selected"])
    _package_program(output_dir)
    (output_dir / "screening.json").write_text(json.dumps(metrics, ensure_ascii=False, indent=2), encoding="utf-8")
    return metrics


def main() -> None:
    parser = argparse.ArgumentParser(description="Problem 3 communication and relay scheduler")
    parser.add_argument("--output", type=Path, default=OUTPUT_DEFAULT)
    parser.add_argument("--sample-step", type=float, default=10.0)
    parser.add_argument("--relay-candidate-step", type=float, default=10.0)
    parser.add_argument("--q2-output", type=Path, default=ROOT / "outputs/problem23/clearance50")
    parser.add_argument("--max-iterations", type=int, default=12)
    parser.add_argument("--time-limit", type=float, default=180.0)
    parser.add_argument("--objective", choices=("sorties", "delay"), default="sorties")
    parser.add_argument("--max-relay-sorties", type=int)
    parser.add_argument("--max-transport-delay", type=float, default=10800.0)
    args = parser.parse_args()
    if (args.sample_step <= 0 or args.relay_candidate_step <= 0
            or args.max_iterations < 1 or args.time_limit <= 0
            or args.max_transport_delay <= 0
            or (args.max_relay_sorties is not None and args.max_relay_sorties < 0)):
        parser.error("sampling steps must be positive")
    result = run(args.output, args.sample_step, args.relay_candidate_step, args.q2_output,
                 args.max_iterations, args.time_limit, args.objective, args.max_relay_sorties,
                 args.max_transport_delay)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if not result["feasible"]:
        raise SystemExit(2)


if __name__ == "__main__":
    main()
