"""Regression checks for frozen-task partitioning and resource accounting."""
import csv
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from final_code.problem4.solver import (
    RESOURCE_KEYS, SITES, _charge_s, atomic_units, evaluate_partition,
    group_attribution, partitions, q3_gate, resource_intervals, run,
)


def write_csv(path, rows, fields=None):
    with path.open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=fields or list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def frozen_fixture(root):
    source = root / 'q3'
    source.mkdir()
    (source / 'screening.json').write_text(json.dumps({
        'feasible': True, 'continuity_certified': True,
        'continuous_communication_validation': {'audit_version': 'exact_dem_interval_v2'},
    }))
    transport, deliveries = [], []
    for index in range(3):
        sites = SITES[index * 5:index * 5 + 5]
        transport.append({
            '架次编号': f'T{index}', '机型编号': 'A', '准备开始时刻（s）': index * 1000,
            '起飞时刻（s）': index * 1000 + 10, '返回O01时刻（s）': index * 1000 + 100,
            '访问服务区顺序': ','.join(sites), '架次能耗（kWh）': 0.2,
            '返航SOC（%）': 80, '逐箱交付数': 5,
        })
        deliveries.extend({'货箱编号': f'{site}-B1', '服务区编号': site,
                           '架次编号': f'T{index}', '质量（kg）': 2} for site in sites)
    write_csv(source / 'transport_inherited_audit.csv', transport)
    write_csv(source / 'box_delivery_audit.csv', deliveries)
    write_csv(source / 'relay_schedule.csv', [], ['relay_sortie', 'covered_sorties'])
    write_csv(source / 'relay_resource_audit.csv', [], ['中继架次编号'])
    write_csv(source / 'communication_audit.csv', [
        {'运输架次编号': row['架次编号'], '保障方式': '直连', '中继架次编号': ''}
        for row in transport
    ])
    return source


class QuestionFourAuditTests(unittest.TestCase):
    def test_default_enumeration_exceeds_previous_limit(self):
        units = [frozenset({f'S{i:03d}'}) for i in range(1, 14)]
        # Stirling S(13, 3) = (3**13 - 3*2**13 + 3) / 6.
        self.assertEqual(sum(1 for _ in partitions(units, 3)), 261625)
        with self.assertRaises(RuntimeError):
            list(partitions(units[:4], 3, max_candidates=2))

    def test_atomic_units_merge_only_shared_transport_sorties(self):
        # Sites sharing a transport sortie cannot be split; sites that merely
        # share a relay sortie are not merged into one indivisible unit.
        transport = [{'架次编号': 'T1', '访问服务区顺序': 'S001,S002'},
                     {'架次编号': 'T2', '访问服务区顺序': 'S003,S004'}]
        units = atomic_units(transport)
        self.assertIn(frozenset({'S001', 'S002'}), units)
        self.assertIn(frozenset({'S003', 'S004'}), units)

    def test_shared_relay_is_fielded_once_per_group_it_serves(self):
        transport = [{'架次编号': 'T1', '访问服务区顺序': 'S001,S002'},
                     {'架次编号': 'T2', '访问服务区顺序': 'S003'}]
        relay = [{'relay_sortie': 'R1', 'covered_sorties': 'T1→T2'}]
        first = frozenset({'S001', 'S002'})
        second = frozenset({'S003'})
        # A partition that splits T1 and T2 stays admissible: no relay mission is
        # split or re-timed, each group simply owns its own aircraft for it.
        own_first, shared_first = group_attribution(first, transport, relay)
        own_second, shared_second = group_attribution(second, transport, relay)
        self.assertEqual(own_first, {'T1', 'R1'})
        self.assertEqual(own_second, {'T2', 'R1'})
        self.assertEqual((shared_first, shared_second), (1, 1))
        # A partition that keeps both sorties together fields R1 only once.
        merged, shared = group_attribution(first | second, transport, relay)
        self.assertEqual(merged, {'T1', 'T2', 'R1'})
        self.assertEqual(shared, 0)

    def test_shared_relay_counts_once_per_group_in_resource_demand(self):
        transport = [{'架次编号': 'T1', '机型编号': 'A', '准备开始时刻（s）': 0,
                      '返回O01时刻（s）': 100, '返航SOC（%）': 100,
                      '访问服务区顺序': 'S001,S002', '架次能耗（kWh）': 0.1},
                     {'架次编号': 'T2', '机型编号': 'A', '准备开始时刻（s）': 0,
                      '返回O01时刻（s）': 100, '返航SOC（%）': 100,
                      '访问服务区顺序': 'S003', '架次能耗（kWh）': 0.1}]
        relay = [{'relay_sortie': 'R1', 'covered_sorties': 'T1→T2',
                  '中继架次编号': 'R1', '准备开始时刻（s）': 0,
                  '建链完成/服务开始（s）': 10, '服务结束时刻（s）': 50,
                  '返回O01时刻（s）': 60, '机体再次可用（s）': 70,
                  '能源组件再次可用（s）': 80, '返航SOC（%）': 100,
                  '架次能耗（kWh）': 0.05}]
        groups = (frozenset({'S001', 'S002'}), frozenset(set(SITES) - {'S001', 'S002'}))
        intervals = resource_intervals(transport, relay)
        with patch('final_code.problem4.solver._battery_full_charge_s', return_value={'A': 100}):
            rows = evaluate_partition(groups, transport, relay, {}, intervals)
        # The overlapping single relay mission forces one airframe in each group,
        # so the partition-level total exceeds the global peak of 1.
        self.assertEqual([row['relay_airframe'] for row in rows], [1, 1])
        self.assertEqual([row['shared_relay_missions'] for row in rows], [1, 1])

    def test_battery_recharge_is_counted_and_frozen_timestamp_checked(self):
        row = {'架次编号': 'T1', '机型编号': 'A', '准备开始时刻（s）': 0,
               '返回O01时刻（s）': 100, '返航SOC（%）': 80}
        with patch('final_code.problem4.solver._battery_full_charge_s', return_value={'A': 100}):
            intervals = resource_intervals([row], [])
            self.assertAlmostEqual(intervals['A_battery'][0][1], 100 + _charge_s(80, 100))
            row['电池充电完成时刻（s）'] = 100
            with self.assertRaisesRegex(ValueError, '充电'):
                resource_intervals([row], [])

    def test_rejects_old_sampled_communication_certificate(self):
        with tempfile.TemporaryDirectory() as temporary:
            source = frozen_fixture(Path(temporary))
            self.assertTrue(q3_gate(source)['ready'])
            screening = json.loads((source / 'screening.json').read_text())
            del screening['continuous_communication_validation']
            (source / 'screening.json').write_text(json.dumps(screening))
            gate = q3_gate(source)
            self.assertFalse(gate['ready'])
            self.assertIn('exact_dem_interval_v2', gate['reason'])

    def test_stock_gap_workload_and_input_conservation(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = frozen_fixture(root)
            original = {path.name: path.read_bytes() for path in source.iterdir()}
            inventory = {key: 0 for key in RESOURCE_KEYS}
            inventory.update(A_airframe=1, A_battery=1)
            with patch('final_code.problem4.solver.load_inventory', return_value=inventory), \
                 patch('final_code.problem4.solver._battery_full_charge_s', return_value={'A': 100}):
                result = run(source, root / 'q4')
            for k, deficit in ((2, 2), (3, 4)):
                solution = result['solutions'][str(k)]
                self.assertTrue(solution['partition_search_complete'])
                self.assertFalse(solution['inventory_feasible'])
                self.assertEqual(solution['score'][0], deficit)
                self.assertEqual(sum(row['cargo_mass_kg'] for row in solution['groups']), 30)
                self.assertEqual(sum(row['box_count'] for row in solution['groups']), 15)
            self.assertEqual(original, {path.name: path.read_bytes() for path in source.iterdir()})
            self.assertEqual(set(result['frozen_input_sha256']), set(original))
            with (root / 'q4/inventory_comparison.csv').open(encoding='utf-8-sig') as stream:
                comparison = list(csv.DictReader(stream))
            self.assertNotIn('group_surplus', comparison[0])
            self.assertNotIn('group_deficits', comparison[0])
            self.assertEqual({row['inventory_deficit'] for row in comparison
                              if row['K'] == '2' and row['resource'] == 'A_airframe'}, {'1'})
            with (root / 'q4/workload_comparison.csv').open(encoding='utf-8-sig') as stream:
                workload = list(csv.DictReader(stream))
            self.assertIn('box_count_cv', workload[0])
            self.assertIn('cargo_mass_kg_cv', workload[0])
            workbook = load_workbook(root / 'q4/problem4_submission.xlsx', read_only=True, data_only=True)
            self.assertEqual(sum(row[0] is not None for row in workbook['Q4_分区配置'].iter_rows(
                min_row=2, values_only=True)), 5)
            workbook.close()


if __name__ == '__main__':
    unittest.main()
