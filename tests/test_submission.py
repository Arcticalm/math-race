"""The six template tables and Q3 transport supplements must stay distinct."""
import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook

from final_code.problem1.solver import ROOT
from final_code.submission import QUESTION_SHEETS, organize_workbooks


def fixture(root):
    for question, names in QUESTION_SHEETS.items():
        folder = root / question
        folder.mkdir()
        book = load_workbook(ROOT / 'docs/结果提交模板.xlsx')
        for name in names:
            book[name].cell(2, 1, f'{question}-{name}')
        book.save(folder / f'problem{question[1:]}_submission.xlsx')
        book.close()


class SubmissionTests(unittest.TestCase):
    def test_keeps_template_headers_and_both_transport_solutions(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            fixture(root)
            counts = organize_workbooks(root)
            combined = load_workbook(root / '结果提交汇总.xlsx', read_only=True)
            template = load_workbook(ROOT / 'docs/结果提交模板.xlsx', read_only=True)
            self.assertEqual(combined.sheetnames[:6], template.sheetnames)
            for name in template.sheetnames:
                self.assertEqual([c.value for c in combined[name][1]], [c.value for c in template[name][1]])
            self.assertEqual(combined['Q2_运输架次'].cell(2, 1).value, 'q2-Q2_运输架次')
            self.assertEqual(combined['Q3_运输架次'].cell(2, 1).value, 'q3-Q2_运输架次')
            self.assertEqual(counts['q3']['Q3_逐箱交付'], 1)
            self.assertTrue((root / 'q3/tables/Q3_运输架次.csv').exists())
            combined.close()
            template.close()
            for question, names in QUESTION_SHEETS.items():
                book = load_workbook(root / question / f'problem{question[1:]}_submission.xlsx', read_only=True)
                self.assertEqual(book.sheetnames, list(names))
                book.close()

    def test_rejects_a_modified_template_header(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            fixture(root)
            path = root / 'q4/problem4_submission.xlsx'
            book = load_workbook(path)
            book['Q4_分区配置'].cell(1, 1, 'wrong')
            book.save(path)
            book.close()
            with self.assertRaisesRegex(ValueError, 'header mismatch'):
                organize_workbooks(root)
            self.assertFalse((root / '结果提交汇总.xlsx').exists())


if __name__ == '__main__':
    unittest.main()
